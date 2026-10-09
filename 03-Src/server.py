"""Astra host. The frozen engine resolves every writable path inside its own copy."""
from __future__ import annotations
import asyncio
from contextlib import asynccontextmanager
import collaboration_lifecycle
import io
import json
import re
import sys
import zipfile
from pathlib import Path

from fastapi import Body, FastAPI, Form, HTTPException, Request
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from starlette.datastructures import UploadFile
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
ENGINE = ROOT / '03-Src/engine'
sys.path.insert(0, str(ENGINE/'backend'))
from app import main as engine
from app import paths
from documents import documents_router
from workspace import workspace_router
from weknora import router as weknora_router
import weknora
from weknora_wiki import router as weknora_wiki_router
from gork_link import router as gork_router
import professional_host
import business_settings
import row_assist
import model_adapter
professional_host.install(engine, paths)
import knowledge_host
knowledge_host.install(paths)

assert paths.PROJECT_ROOT == ENGINE, 'Refuse to use an external business engine'
@asynccontextmanager
async def host_lifespan(_app):
    from app import feishu_app_bot
    await asyncio.to_thread(collaboration_lifecycle.resume_receivers, feishu_app_bot,
                            bool(professional_host.config().get('external_enabled')))
    watchdog=asyncio.create_task(collaboration_lifecycle.monitor_receivers(
        feishu_app_bot, lambda: professional_host.config().get('external_enabled')))
    try:
        yield
    finally:
        watchdog.cancel()
        try:await watchdog
        except asyncio.CancelledError:pass

app = FastAPI(title='智算Astra', version='0.9.1', docs_url=None, redoc_url=None, lifespan=host_lifespan)
import hashlib
HOST_SOURCE_SHA256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
HOST_SOURCES={name:hashlib.sha256((ROOT/'03-Src'/name).read_bytes()).hexdigest() for name in ['server.py','collaboration_lifecycle.py','weknora.py','weknora_wiki.py','professional_host.py','knowledge_host.py','business_settings.py','row_assist.py','model_adapter.py','gork_notifications.py','gork_config.py','gork_stages.py','vendor/gork-output/external_client.py']}
UUID = re.compile(r'^[a-f0-9]{32}$')
MAX_UPLOAD = 30 * 1024 * 1024
write_lock = asyncio.Lock()

READ = {'/api/project-default-settings', '/api/professional-skills', '/api/risk/summary',
        '/api/projects/dashboard', '/api/projects', '/api/projects/{project_id}',
        '/api/projects/{project_id}/runs', '/api/projects/{project_id}/artifacts',
        '/api/projects/{project_id}/artifacts/{artifact_id}/download',
        '/api/standard-trace', '/api/process/{job_id}/batch-match-status',
        '/api/download/{job_id}/{kind}', '/api/tasks/{task_id}', '/api/tasks/{task_id}/timeline', '/api/projects/{project_id}/tasks',
        '/api/knowledge/libraries', '/api/knowledge-articles', '/api/knowledge-articles/graph', '/api/knowledge-articles/{article_id}'}
WRITE = {'/api/inspect', '/api/process', '/api/process/batch-match', '/api/preview/refresh',
         '/api/preview/cell', '/api/preview/recalculate', '/api/experience-warnings/run',
         '/api/fill-assist/candidates', '/api/risk-report', '/api/knowledge/search'}
for route in engine.app.routes:
    if (getattr(route, 'path', '') in (READ | WRITE) or professional_host.accepts(getattr(route,'path',''))) and getattr(route,'path','') not in {'/api/collaboration/external-dispatch/web-review','/api/collaboration/external-dispatch/tasks/{task_id}/retry','/api/risk-report','/api/preview/cell','/api/preview/recalculate','/api/workload-capture/apply-to-current','/api/workload-capture/run'}:
        app.router.routes.append(route)

@app.post('/api/collaboration/external-dispatch/web-review')
def connected_web_review(payload: dict = Body(...)):
    professional_host.require_review_connection(str(payload.get('platform_profile_id') or '').strip())
    return engine.create_web_result_review(payload)

@app.post('/api/collaboration/external-dispatch/tasks/{task_id}/retry')
def connected_review_retry(task_id: str):
    task=engine.external_task_dispatch.ExternalDispatchStore().get_task(task_id)
    if not task:raise HTTPException(404,'未找到外部派发任务')
    professional_host.require_review_connection(str(task.get('platform_profile_id') or ''))
    return engine.retry_external_dispatch_task(task_id)

def checked_job(job_id):
    if not UUID.fullmatch(str(job_id or '')):
        raise HTTPException(400, '任务编号格式无效')
    path=paths.RUNTIME_DIR / job_id
    if not (path/'process-state.json').is_file():
        raise HTTPException(404, '独立工作区中未找到该任务')
    return path

def document_report_path(job_id):
    folder = checked_job(job_id)
    state = engine._load_process_state(folder)
    file = engine._state_path(folder, state, 'output_report', required=False)
    if not file or not file.is_file():
        raise HTTPException(404, '当前任务还没有 Word 报告')
    file.resolve().relative_to(folder.resolve())
    return file

app.include_router(documents_router(lambda: paths.RUNTIME_DIR, document_report_path))
app.include_router(workspace_router(lambda: paths.RUNTIME_DIR, checked_job))
app.include_router(weknora_router)
app.include_router(weknora_wiki_router)
app.include_router(gork_router)
import gork_config
import gork_notifications
app.include_router(gork_config.router)
app.include_router(professional_host.router)
app.include_router(business_settings.router)
app.include_router(row_assist.router)
# Original professional UI uses the same verified Astra WeKnora adapter.
app.add_api_route('/api/knowledge/capabilities',weknora.capabilities,methods=['GET'])
app.add_api_route('/api/knowledge/ask-stream',weknora.ask_stream,methods=['POST'])
app.add_api_route('/api/knowledge/requests/{request_id}/cancel',weknora.cancel,methods=['POST'])
app.add_api_route('/api/knowledge/citations/{citation_id}',weknora.citation,methods=['GET'])

@app.middleware('http')
async def boundaries(request: Request, call_next):
    if request.url.path.startswith('/api/'):
        origin=request.headers.get('origin')
        if origin and origin not in {'http://127.0.0.1:5189', 'http://localhost:5189'}:
            return JSONResponse({'detail':'当前接口仅供本地 Astra 工作区访问'},403)
        # Professional APIs stay same-origin. External sending is an explicit independent setting.
        if request.method in {'POST','PUT','PATCH'}:
            chunks=[];size=0
            async for chunk in request.stream():
                size+=len(chunk)
                if size>MAX_UPLOAD:return JSONResponse({'detail':'文件或请求超过 30 MB 上限'},413)
                chunks.append(chunk)
            request._body=b''.join(chunks)
            try:
                if 'application/json' in request.headers.get('content-type',''):
                    payload=await request.json()
                    if not isinstance(payload,dict):raise ValueError('请求必须是对象')
                else:
                    payload=dict(await request.form())
                if payload.get('job_id') and not request.url.path.startswith('/api/astra/weknora/'):
                    checked_job(str(payload['job_id']))
                for value in payload.values():
                    if isinstance(value,UploadFile):
                        if request.url.path == '/api/astra/workspace/assets':
                            continue  # Dedicated immutable-file validation in workspace.py.
                        name=value.filename or ''
                        if not name.lower().endswith('.xlsx') or any(c in name for c in '/\\:') or name.startswith('.'):
                            raise ValueError('请选择名称有效的 .xlsx 文件')
                        data=await value.read();await value.seek(0)
                        try:
                            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                                if sum(i.file_size for i in archive.infolist())>150*1024*1024:
                                    raise ValueError('工作簿解压体积超过 150 MB 上限')
                        except zipfile.BadZipFile:raise ValueError('文件不是有效的 Excel 工作簿')
                if request.url.path=='/api/process':
                    if payload.get('matching_mode','mechanical') not in {'mechanical','ai_enhanced'}:
                        raise ValueError('编制模式无效')
                    if str(payload.get('defer_matching','')).lower()!='true':
                        raise ValueError('请先转换，再明确发起批量匹配')
                if request.url.path=='/api/risk-report':
                    state=engine._load_process_state(checked_job(str(payload.get('job_id'))))
                    if not engine._is_demo_report_task(state) and not professional_host.model_key(str(payload.get('provider') or professional_host.model_preference().get('provider','deepseek'))):
                        return JSONResponse({'detail':'当前未配置模型。正式 Word 匹配报告已生成；风险润色待模型接通。'},409)
            except HTTPException as exc:return JSONResponse({'detail':exc.detail},exc.status_code)
            except (ValueError,json.JSONDecodeError) as exc:return JSONResponse({'detail':str(exc)},400)
            if request.url.path.startswith('/api/collaboration/') and not professional_host.config().get('external_enabled') and any(token in request.url.path for token in ('/notify','/test','/retry','/settings')):
                return JSONResponse({'detail':'请先在 Astra 连接设置中明确启用外部协同，再配置平台。当前未发送消息。'},409)
            if request.url.path in {'/api/collaboration/external-dispatch/tasks','/api/collaboration/external-dispatch/web-review'} and not professional_host.config().get('external_enabled'):
                return JSONResponse({'detail':'尚未启用外部协同，未发起派发'},409)
            if request.url.path in {'/api/astra/row-assist/candidates','/api/astra/row-assist/details','/api/astra/row-assist/ask'} or request.url.path.startswith(('/api/astra/workspace/', '/api/astra/weknora/', '/api/astra/gork/notifications', '/api/knowledge/ask-stream', '/api/knowledge/requests/')) or request.url.path.endswith('/ai-matching/cancel'):

                try:
                    response = await call_next(request)
                    response.headers['X-Astra-Workspace'] = 'isolated'
                    return response
                finally:
                    await request.close()
            if write_lock.locked():return JSONResponse({'detail':'当前操作仍在执行，请完成后再提交'},409)
            async with write_lock:
                try:
                    response = await call_next(request)
                    response.headers['X-Astra-Workspace'] = 'isolated'
                    return response
                finally:
                    await request.close()
        job=request.query_params.get('job_id')
        if job:
            try:checked_job(job)
            except HTTPException as exc:return JSONResponse({'detail':exc.detail},exc.status_code)
        parts=request.url.path.split('/')
        if len(parts)>3 and parts[2] in ('download','process'):
            try:checked_job(parts[3])
            except HTTPException as exc:return JSONResponse({'detail':exc.detail},exc.status_code)
    response=await call_next(request)
    response.headers['X-Astra-Workspace']='isolated'
    if request.url.path.startswith('/system-architecture/'):
        response.headers['Content-Security-Policy'] = "default-src 'none'; script-src 'unsafe-inline'; style-src 'unsafe-inline'; img-src 'self' data: blob:; font-src 'self' data:; connect-src 'none'; object-src 'none'; base-uri 'none'; form-action 'none'"
    return response

@app.get('/api/health')
def health():
    import hashlib, os
    package=json.loads((ROOT/'package.json').read_text(encoding='utf-8'))
    build_file=ROOT/'04-output/astra-build.json'
    try: build=json.loads(build_file.read_text(encoding='utf-8'))
    except (OSError,ValueError):build=None
    provider=str(professional_host.model_preference().get('provider','deepseek'))
    return {'status':'ok','service':'zhisuan-astra','version':app.version,'ui_version':package['version'],
            'host_sources':HOST_SOURCES,'host_source_sha256':HOST_SOURCE_SHA256,'build':build,'pid':os.getpid(),'workspace':str(ROOT),
            'runtime':str(paths.RUNTIME_DIR),'capabilities':['commands-v1','weknora-wiki-read-v1','batch-match-recovery-v1','row-assist-v1','business-settings-v1','gork-stage-notifications-v1','gork-stage-debug-v1'],
            'engine_version':engine.APP_VERSION,'isolated':True,'mode':'mechanical+ai_enhanced',
            'model_connected':False,'model_configured':bool(professional_host.model_key(provider)),
            'model_provider':provider,'external_sending':bool(professional_host.config().get('external_enabled',False))}

@app.get('/api/astra/sample')
def sample():
    directory=ENGINE/'08-测试与演示/测试-V3.1-260802'
    file=next(directory.glob('*测试输入*演示*.xlsx'))
    return FileResponse(file,filename=file.name)

@app.post('/api/astra/knowledge/ask')
async def knowledge_answer(payload: dict = Body(...)):
    # The frozen async handler performs blocking retrieval / model I/O internally.
    # Keep it off the ASGI loop so the independent WeKnora stream can progress.
    return await asyncio.to_thread(_knowledge_answer_sync, payload)

def _knowledge_answer_sync(payload: dict):
    preference=professional_host.model_preference()
    payload={'provider':preference.get('provider'),'model':preference.get('model'),'base_url':preference.get('baseUrl'),**payload}
    question = str(payload.get('question') or '').strip()
    if not question or len(question) > 10000:
        raise HTTPException(400, '请填写 1 至 10000 字的问题')
    # Original curated answers are explicitly identified. Never call an unconfigured provider.
    cleaned, _ = engine.strip_force_knowledge_prefix(question)
    if engine.get_demo_answer(cleaned) or professional_host.model_key(str(payload.get('provider','deepseek'))):
        with model_adapter.knowledge_request(preference):
            return asyncio.run(engine.knowledge_ask(payload))
    data = asyncio.run(engine.knowledge_search({k: v for k, v in payload.items()
        if k in {'question', 'library_ids', 'retrieval_mode', 'project_key', 'job_id', 'limit'}}))
    return {**data, 'sources': data['results'], 'generated_by_model': False, 'model_configured': False,
            'answer_mode': 'evidence_only',
            'degradation_reason': '所选模型未配置独立凭证，尚未生成问答答案',
            'answer': '尚未生成答案：所选模型未配置独立凭证。\n\n已检索到原文依据，可先核对引用；配置模型后重新提问。' if data['evidence_found'] else '尚未生成答案：所选模型未配置独立凭证，当前范围也未检索到可引用依据。请先配置模型，并核对知识库范围。'}

app.add_api_route('/api/knowledge/ask',knowledge_answer,methods=['POST'])

@app.post('/api/astra/row-assist/ask')
async def row_answer(payload:dict=Body(...)):
    context=row_assist.validate(payload)
    original=await row_assist.checked_candidates(payload)
    answer=await knowledge_answer({'question':payload.get('question'),'job_id':context['job_id'],
        'context_type':'row','row_context':{**original['context'],'candidate_recommendations':row_assist.recommendation_shortlist(original['candidates'])},
        'retrieval_mode':'classic'})
    row_assist.validate(payload)  # A late answer must not describe a newer artifact as the captured row.
    return answer

@app.get('/api/astra/collaboration/status')
def collaboration_status():
    status=engine.get_feishu_app_bot_status()
    platforms=[]
    for profile in status.get('profiles',[]):
        health=profile.get('health') or {}
        verified=bool(profile.get('running') and health.get('fresh') and health.get('connected') and health.get('identity_ready') and health.get('backend_ok') and health.get('healthy'))
        platforms.append({'id':'weact' if profile.get('profile_id')=='weact_cost' else 'feishu',
            'name':'WeAct' if profile.get('profile_id')=='weact_cost' else '飞书',
            'configured':bool(profile.get('configured')),'enabled':bool(profile.get('enabled')),
            'connected':verified,'verified':verified,'delivery_verified':False})
    return {'external_sending':bool(professional_host.config().get('external_enabled')), 'platforms':platforms,
        'note':'连接状态来自当前健康快照，真实送达以每次派发回执为准。'}

@app.get('/api/astra/runs')
def runs():
    result=[]
    for file in paths.RUNTIME_DIR.glob('*/process-state.json'):
        if not UUID.fullmatch(file.parent.name):continue
        try:
            state=json.loads(file.read_text(encoding='utf-8'))
            summary=state.get('summary',{})
            result.append({'job_id':file.parent.name,'name':state.get('project_name') or state.get('input_filename'),
                'input_filename':state.get('input_filename'),'updated_at':state.get('updated_at'),
                'status':summary.get('matching_status'),'rows':summary.get('total_data_rows',0)})
        except (OSError,ValueError):continue
    return {'items':sorted(result,key=lambda r:r['updated_at'] or '',reverse=True)[:60]}

@app.get('/api/astra/runs/{job_id}')
def restore(job_id:str):
    directory=checked_job(job_id)
    state=engine._load_process_state(directory)
    downloads={kind:f'/api/download/{job_id}/{kind}' for kind,key in [('excel','output_excel'),('report','output_report')]
               if state.get(key) and (directory/state[key]).is_file()}
    return engine._attach_ai_matching(engine._attach_job_skill({'job_id':job_id,'summary':state['summary'],'downloads':downloads,'skill_snapshot':state.get('skill_snapshot'),
            'input_filename':state.get('input_filename'),'project_name':state.get('project_name'),
            'task_tracking':{'task':state.get('business_task',{})},
            'needs_recalculate':read_edit_state(directory),'process_options':state.get('process_options',{}),'artifact_version':state.get('artifact_version',1),'business_task':state.get('business_task'),
            'project_tracking':state.get('project_tracking')},directory,state),directory,state)


@app.post('/api/projects/backfill')
def dashboard_backfill(include_unclassified_legacy: bool = False):
    # The original aggregation contract, bounded to Astra's own business runtime.
    # Never scan the original project or external collaboration runtime.
    return engine.backfill_project_ledger(
        engine._project_ledger(), runtime_root=paths.RUNTIME_DIR,
        include_unclassified_legacy=include_unclassified_legacy)


@app.get('/api/projects/{project_id}/runs/{run_id}/result')
def dashboard_result(project_id: str, run_id: str):
    try:
        run = engine._project_ledger().get_run(project_id, run_id)
    except engine.ProjectNotFoundError as exc:
        raise HTTPException(404, '独立工作区中未找到该项目任务') from exc
    if not run.get('job_id'):
        raise HTTPException(409, '该任务暂无可恢复的网页成果')
    # Reuse Astra restore, retaining the real manual-edit / pending-recalculation state.
    result = restore(str(run['job_id']))
    return {**result, 'project_tracking': {'project_id': project_id, 'run_id': run_id}}

def read_edit_state(directory:Path):
    file=directory/'astra-edit-state.json'
    if not file.exists():return False
    return bool(json.loads(file.read_text(encoding='utf-8')).get('needs_recalculate'))

@app.post('/api/preview/cell')
async def edit_cell(payload:dict=Body(...)):
    directory=checked_job(str(payload.get('job_id')))
    # Adapter metadata remains separate from the unchanged original engine state.
    result=await engine.update_preview_cell(payload)
    needed=bool(result.get('needs_recalculate')) or not bool(result.get('formula_recalculated'))
    engine.atomic_json(directory/'astra-edit-state.json',{'needs_recalculate':needed})
    return {**result,'needs_recalculate':needed}

@app.post('/api/risk-report')
async def risk_report(job_id:str=Form(...),provider:str|None=Form(default=None),model:str|None=Form(default=None),base_url:str|None=Form(default=None)):
    preference=professional_host.model_preference()
    return await engine.generate_risk_report(job_id=job_id,provider=provider or preference['provider'],model=model or preference['model'],base_url=base_url or preference['baseUrl'])

@app.post('/api/fill-assist/confirm')
async def confirm_candidate(payload:dict=Body(...)):
    directory=checked_job(str(payload.get('job_id')))
    result=await engine.confirm_fill_assist(payload)
    needed=bool(result.get('needs_recalculate')) or not bool(result.get('formula_recalculated'))
    engine.atomic_json(directory/'astra-edit-state.json',{'needs_recalculate':needed})
    return {**result,'needs_recalculate':needed}

# Preserve the original multipart contract and add Astra's durable edit metadata.
import functools
@functools.wraps(engine.apply_workload_capture_to_current)
async def apply_workload(**kwargs):
    result=await engine.apply_workload_capture_to_current(**kwargs)
    directory=checked_job(str(kwargs.get('job_id')))
    engine.atomic_json(directory/'astra-edit-state.json',{'needs_recalculate':not bool(result.get('formula_recalculated'))})
    return result
app.add_api_route('/api/workload-capture/apply-to-current',apply_workload,methods=['POST'])

@functools.wraps(engine.run_workload_capture)
async def capture_workload_pair(**kwargs):
    # Two uploads may share the same filename. Preserve both independent inputs.
    source=kwargs['workload_file'];target=kwargs['target_file']
    if source.filename.casefold()==target.filename.casefold():
        source.filename='工作量源表-'+source.filename
        target.filename='控制价目标表-'+target.filename
    return await engine.run_workload_capture(**kwargs)
app.add_api_route('/api/workload-capture/run',capture_workload_pair,methods=['POST'])

@app.post('/api/preview/recalculate')
async def recalculate(payload:dict=Body(...)):
    directory=checked_job(str(payload.get('job_id')))
    result=await engine.recalculate_preview_workbook(payload)
    needed=not bool(result.get('formula_recalculated'))
    engine.atomic_json(directory/'astra-edit-state.json',{'needs_recalculate':needed})
    return {**result,'needs_recalculate':needed}

@app.get('/api/astra/report/{job_id}')
def report(job_id:str):
    directory=checked_job(job_id)
    state=engine._load_process_state(directory)
    file=engine._state_path(directory,state,'output_report',required=False)
    if not file or not file.is_file():raise HTTPException(404,'请先完成批量匹配，生成正式报告')
    file.resolve().relative_to(directory.resolve())
    document=Document(file)
    return {'filename':file.name,'paragraphs':[p.text for p in document.paragraphs if p.text.strip()],
            'tables':[[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables],
            'note':'读取当前实际 DOCX 的正文与表格；此视图不模拟 Word 分页。完整排版请打开下载文件。'}

@app.get('/api/astra/digital/status')
async def digital_status():
    import httpx
    try:
        async with httpx.AsyncClient(timeout=2,trust_env=False) as client:
            reply=await client.get('http://127.0.0.1:5175/')
        return {'available':reply.status_code==200}
    except httpx.HTTPError:return {'available':False}

gork_notifications.install(app, engine, paths.RUNTIME_DIR)
if (ROOT/'04-output').exists():
    app.mount('/',StaticFiles(directory=ROOT/'04-output',html=True),name='studio')
