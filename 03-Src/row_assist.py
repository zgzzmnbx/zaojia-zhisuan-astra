"""Version-bound row review. Candidates are revalidated before an explicit single-cell write."""
import hashlib
import hmac
import json
import secrets
import math
from statistics import mean, median
from fastapi import APIRouter, Body, HTTPException
from workspace import Workspace, identifier

router=APIRouter(prefix='/api/astra/row-assist')
_secret=secrets.token_bytes(32)
def sign(context):
    return hmac.new(_secret,json.dumps(context,sort_keys=True,ensure_ascii=False,default=str).encode(),hashlib.sha256).hexdigest()

def candidate_digest(items):
    return hashlib.sha256(json.dumps(items,sort_keys=True,ensure_ascii=False,default=str).encode()).hexdigest()

def recommendation_shortlist(items):
    from app.fill_assist import _number
    result=[];seen=set()
    for item in items:
        value=_number(item.get('value'))
        if value is None or not math.isfinite(value) or value in seen:continue
        seen.add(value)
        result.append({**item,'supporting_sources':[c for c in items if _number(c.get('value'))==value]})
        if len(result)==3:break
    return result

async def checked_candidates(payload):
    from app import main as engine
    context=validate(payload)
    data=await field_candidates(context)
    if context.get('candidate_digest')!=candidate_digest(data['candidates']):
        raise HTTPException(409,'计价依据已更新，请重新读取候选后确认')
    validate(payload)
    return data


FIELDS={'实物工作费调整系数':'physical_adjustment','技术工作费调整系数':'technical_adjustment'}

async def field_candidates(payload):
    from app import main as engine, paths
    from app.ai_matching.orchestrator import _candidate_snapshot
    from app.fill_assist import _read_pool_records, _number, build_fill_assist_candidates
    from app.experience_warning import ExperiencePoolMatcher
    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string
    data=await engine.fill_assist_candidates(payload)
    requested=str(payload.get('target_header') or '')
    metric=next((name for name in FIELDS if name in requested),None)
    context=data['context']
    state=engine._load_process_state(paths.RUNTIME_DIR/payload['job_id'])
    runtime=engine._job_skill_runtime_context(paths.RUNTIME_DIR/payload['job_id'],state)
    fill=engine._build_skill_fill_engine(runtime)
    # Use the exact batch-engine input normalization (e.g. 成图比例 is a layout marker).
    context['row']={key:fill._clean_input_value(value) for key,value in context['row'].items()}
    if not metric:
        data['candidates']=build_fill_assist_candidates(context['row'],knowledge_base=fill.knowledge_base,pool_path=runtime.risk_profile.get('experiencePool'))
        return data
    book=load_workbook(engine._state_path(paths.RUNTIME_DIR/payload['job_id'],state,'output_excel'),data_only=True)
    try:
        sheet=book[context['sheet_name']]
        merged=fill._build_merged_value_map(sheet)
        raw=fill._read_row_values(sheet,context['excel_row'],merged)
        headers=[merged.get((context['header_row'],c),sheet.cell(context['header_row'],c).value) for c in range(1,sheet.max_column+1)]
        options=state.get('process_options') or {}
        config=next((c for c in options.get('sheet_configs',[]) if c.get('sheet_name')==context['sheet_name']),{})
        mapping=config.get('column_mapping') or options.get('column_mapping') or {}
        mapped=mapping.get('输出-'+metric)
        columns=[column_index_from_string(str(mapped))-1] if mapped else [i for i,h in enumerate(headers) if str(h or '').replace(' ','').strip()==requested.replace(' ','').strip()]
        if not columns:columns=[i for i,h in enumerate(headers) if metric in str(h or '').replace(' ','') and not any(x in str(h) for x in ('说明','预警','经验','批注','抓取'))]
        if len(columns)!=1:raise HTTPException(400,'当前工作表的系数列缺失或不唯一，请先核对字段映射')
        column=columns[0]
        if not 0<=column<len(headers) or metric not in str(headers[column]):raise HTTPException(400,'系数字段映射与实际表头不一致，请先核对映射')
        context.update(target_column=column+1,target_header=str(headers[column]),current_value=raw[column])
        values={**context['row'],**{str(h):v for h,v in zip(headers,raw) if h}}
    finally:book.close()
    field=FIELDS[metric]
    trusted=_candidate_snapshot({'values':values,'sheet_name':context['sheet_name'],'raw_values':raw},runtime,knowledge_base=fill.knowledge_base,adjustment=fill.adjustment_engine)[field]
    trusted.sort(key=lambda c:0 if c.get('rule_id') else 1)
    for item in trusted:
        item['source_label']='匹配规则' if item.get('rule_id') else '计价库经验层'
        item['reason']=str(item.get('note') or item.get('basis') or '符合当前要素与单位条件')
        item['risk_tips']=[] if item.get('rule_id') else ['计价库经验提示，需核对本项目适用条件。']
    pool=runtime.risk_profile.get('experiencePool')
    if pool and pool.exists():
        mode,description,records=ExperiencePoolMatcher(_read_pool_records(pool)).lookup(context['row'])
        groups={}
        for record in records:
            value=_number(record.get(metric))
            if value is not None and math.isfinite(value):groups.setdefault(value,[]).append(record)
        for value,samples in sorted(groups.items(),key=lambda item:(-len(item[1]),item[0])):
            trusted.append({'id':f'pool-{field}-{value}','field':field,'value':value,'source':'experience_pool','source_label':'经验池同类样本','sample_count':len(samples),'reason':f'{description or mode}，此系数出现 {len(samples)} 次。','basis':f'经验池 {pool.name}','risk_tips':['历史系数按出现次数排序，需核对当前适用条件。']})
    data['candidates']=trusted
    return data
def owner(payload):
    from app import paths
    conversation=str(payload.get('conversation_id') or '')
    identifier(payload.get('job_id'))
    with Workspace(paths.RUNTIME_DIR/'astra-workspace').connect() as db:
        if not db.execute('SELECT id FROM conversations WHERE id=?',(conversation,)).fetchone():raise HTTPException(404,'当前对话不存在')
        bound=db.execute('SELECT conversation_id FROM conversation_jobs WHERE job_id=?',(payload.get('job_id'),)).fetchone()
        if not bound or bound['conversation_id']!=conversation:raise HTTPException(409,'当前成果不属于该对话，请保存任务关联后重试')
def snapshot(payload):
    from app import main as engine, paths
    owner(payload)
    folder=paths.RUNTIME_DIR/payload['job_id']
    state=engine._load_process_state(folder)
    file=engine._state_path(folder,state,'output_excel')
    return hashlib.sha256(file.read_bytes()).hexdigest(),state
def artifact_version(state):
    return int(state.get('artifact_version') or (state.get('business_task') or {}).get('artifact_version') or 1)

def validate(payload):
    context=payload.get('row_context')
    if not isinstance(context,dict) or not hmac.compare_digest(str(payload.get('token') or ''),sign(context)):raise HTTPException(409,'行上下文已失效，请重新打开该行')
    sha,state=snapshot(context)
    if sha!=context.get('artifact_sha256') or artifact_version(state)!=context.get('artifact_version'):raise HTTPException(409,'成果已更新，旧候选不能覆盖当前版本，请重新打开该行')
    return context
@router.post('/candidates')
async def candidates(payload:dict=Body(...)):
    from app import main as engine, paths
    sha,state=snapshot(payload)
    if payload.get('expected_artifact_sha256') and payload['expected_artifact_sha256']!=sha:raise HTTPException(409,'成果已更新，请从主表重新打开该行')
    data=await field_candidates(payload)
    context=data['context']
    if context['excel_row'] <= context['header_row'] or not str(context['row'].get('单位') or '').strip():
        raise HTTPException(400,'请选择有计价单位的业务数据行，表头和空行不能填价')
    if not context.get('target_column'):raise HTTPException(400,'当前行未找到可填价字段')
    if not any(word in str(context['target_header']) for word in ('基价','单价',*FIELDS)):raise HTTPException(400,'请选择基价、单价或工作费调整系数字段')
    operation_file=paths.RUNTIME_DIR/payload['job_id']/'batch-match-status.json'
    operation=json.loads(operation_file.read_text(encoding='utf-8')) if operation_file.exists() else {}
    row={'job_id':payload['job_id'],'conversation_id':payload['conversation_id'],'context':str(payload.get('context') or ''),'sheet_name':context['sheet_name'],'row_number':context['excel_row'],'column_number':context['target_column'],'target_header':context['target_header'],'artifact_sha256':sha,'artifact_version':artifact_version(state),'task_id':str((state.get('business_task') or {}).get('task_id') or ''),'run_id':str((state.get('project_tracking') or {}).get('run_id') or (state.get('ai_matching') or {}).get('run_id') or ''),'attempt_id':str(operation.get('attempt') or ''),'current_value':context.get('current_value')}
    row['candidate_digest']=candidate_digest(data['candidates'])
    if snapshot(payload)[0]!=sha:raise HTTPException(409,'成果正在更新，请稍后重新打开该行')
    data.update(row_context=row,token=sign(row))
    return data


def distribution(values):
    values=sorted(float(v) for v in values if isinstance(v,(int,float)) and not isinstance(v,bool) and math.isfinite(v))
    if not values:return {'count':0,'bins':[]}
    low,high=values[0],values[-1]
    count=min(10,max(1,math.ceil(math.sqrt(len(values))))) if high>low else 1
    width=(high-low)/count if high>low else 0
    bins=[{'min':low+i*width,'max':low+(i+1)*width,'count':0} for i in range(count)]
    for value in values:bins[min(count-1,int((value-low)/width)) if width else 0]['count']+=1
    return {'count':len(values),'min':low,'max':high,'mean':mean(values),'median':median(values),'bins':bins}


@router.post('/details')
async def details(payload:dict=Body(...)):
    """Read the same matching scope as the frozen engine, never invent sample values."""
    from app import main as engine, paths
    from app.fill_assist import _knowledge_similar_candidates, _read_pool_records, _number, _candidate_sort_key
    from app.experience_warning import ExperiencePoolMatcher, PRICE_METRIC
    from app.knowledge_base import KnowledgeBase
    context=validate(payload)
    original=await checked_candidates(payload)
    row=original['context']['row']
    folder=paths.RUNTIME_DIR/context['job_id']
    state=engine._load_process_state(folder)
    runtime=engine._job_skill_runtime_context(folder,state)
    metric=next((name for name in FIELDS if name in context['target_header']),PRICE_METRIC)
    rules=sorted(_knowledge_similar_candidates(row,KnowledgeBase.from_excel(runtime.knowledge_base_path)),key=_candidate_sort_key) if metric==PRICE_METRIC else [c for c in original['candidates'] if c.get('source')!='experience_pool']
    pool=runtime.risk_profile.get('experiencePool',engine._resolve_experience_pool_path())
    matched=[];mode='';description=''
    if pool and pool.exists():mode,description,matched=ExperiencePoolMatcher(_read_pool_records(pool)).lookup(row)
    samples=[]
    for record in matched:
        value=_number(record.get(metric))
        if value is not None and math.isfinite(value):
            samples.append({'value':value,'source_file':record.get('来源文件',''),'source_sheet':record.get('来源sheet',''),'source_row':record.get('来源行',''),'pool_row':record.get('_excel_row','')})
    samples.sort(key=lambda r:r['value'])
    validate(payload)
    return {'rules':{'items':rules[:200],'total':len(rules),'distribution':distribution([_number(r['value']) for r in rules])},
            'experience':{'items':samples[:200],'total':len(samples),'match_mode':mode,'match_detail':description,'distribution':distribution([r['value'] for r in samples])},
            'trace':original.get('trace',[]),'context':original['context']}
@router.post('/confirm')
async def confirm(payload:dict=Body(...)):
    from app import main as engine, paths
    context=validate(payload)
    original=await checked_candidates(payload)
    candidate=next((item for item in original['candidates'] if item['id']==payload.get('candidate_id')),None)
    if candidate is None:raise HTTPException(409,'候选已失效，请重新读取依据')
    result=await engine.confirm_fill_assist({**context,'candidate':candidate,'note':str(payload.get('note') or ''),'header_rows':payload.get('header_rows',{}),'actor':'Astra 本机用户','recalculate':False})
    needed=bool(result.get('needs_recalculate')) or not bool(result.get('formula_recalculated'))
    engine.atomic_json(paths.RUNTIME_DIR/context['job_id']/'astra-edit-state.json',{'needs_recalculate':needed})
    return {**result,'needs_recalculate':needed}
