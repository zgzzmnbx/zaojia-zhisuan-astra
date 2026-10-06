"""Project location recognition and spatial ledger support.

This module deliberately keeps location recognition separate from the pricing
engine.  It only produces a reviewable suggestion or stores a human-confirmed
spatial profile; it never writes prices, adjustment factors, or workbook rows.
"""

from __future__ import annotations

import json
import hashlib
import re
import sqlite3
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .project_ledger import ProjectNotFoundError, validate_project_id


SPATIAL_TYPES = {"POINT", "AREA", "CORRIDOR"}
SPATIAL_STATUSES = {"unlocated", "proposed", "confirmed", "conflict", "rejected"}
SPATIAL_ACTIONS = {"propose", "confirm", "correct", "reject", "clear", "migrate"}
DEFAULT_ADMIN_DICTIONARY_PATH = (
    Path(__file__).resolve().parent / "assets" / "admin_divisions" / "approved-admin-v1.json"
)
SOURCE_VERSION = "admin-display-centroid-v1"
MAP_PROVIDER = "antv-l7"
MAP_RUNTIME_MODE = "l7_test"
MAX_WORKBOOK_SHEETS = 12
MAX_WORKBOOK_ROWS = 120
MAX_WORKBOOK_COLUMNS = 80
MAX_TEXT_LENGTH = 240
MAX_EVIDENCE_ITEMS = 40
IGNORED_RECOGNITION_VALUES = {
    "",
    "-",
    "--",
    "暂无",
    "无",
    "待填写",
    "待补充",
    "模板",
    "示例",
    "样例",
    "demo",
    "placeholder",
    "n/a",
    "na",
    "null",
    "none",
}


class ProjectSpatialError(RuntimeError):
    """Base error for the isolated spatial module."""


class ProjectSpatialNotFound(ProjectSpatialError):
    pass


class ProjectSpatialValidationError(ProjectSpatialError):
    pass


class ProjectSpatialConflict(ProjectSpatialError):
    pass


def now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def _date_only(value: object) -> date | None:
    try:
        return datetime.fromisoformat(str(value or "").replace("Z", "+00:00")).date()
    except (TypeError, ValueError):
        return None


def clean_text(value: object, limit: int = MAX_TEXT_LENGTH) -> str:
    return str(value or "").replace("\x00", "").strip()[:limit]


def map_runtime_enabled(dictionary: dict[str, Any]) -> bool:
    """The test project may render L7 without a formal map-approval gate."""

    return (
        clean_text(dictionary.get("map_provider"), 40).casefold() == MAP_PROVIDER
        and clean_text(dictionary.get("map_runtime_mode"), 40).casefold() == MAP_RUNTIME_MODE
    )


def normalize_text(value: object) -> str:
    """Normalize only for deterministic lookup; preserve the original evidence."""

    text = unicodedata.normalize("NFKC", clean_text(value, 600)).casefold()
    return re.sub(r"[\s\-_—–·,，。；;:：/\\|()（）【】\[\]{}<>《》]", "", text)


def _safe_json(value: object, *, limit: int = MAX_TEXT_LENGTH) -> object:
    if isinstance(value, dict):
        return {clean_text(key, 80): _safe_json(item, limit=limit) for key, item in list(value.items())[:50]}
    if isinstance(value, list):
        return [_safe_json(item, limit=limit) for item in value[:50]]
    if isinstance(value, (str, int, float, bool)) or value is None:
        return clean_text(value, limit) if isinstance(value, str) else value
    return clean_text(value, limit)


def load_admin_dictionary(path: Path | None = None) -> dict[str, Any]:
    dictionary_path = Path(path or DEFAULT_ADMIN_DICTIONARY_PATH)
    try:
        payload = json.loads(dictionary_path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise ProjectSpatialError(f"地点字典不存在：{dictionary_path}") from exc
    except (OSError, json.JSONDecodeError) as exc:
        raise ProjectSpatialError(f"地点字典不可读：{dictionary_path}") from exc
    if not isinstance(payload, dict) or not isinstance(payload.get("records"), list):
        raise ProjectSpatialError("地点字典结构无效")
    records: list[dict[str, Any]] = []
    seen_ids: set[str] = set()
    for raw in payload["records"]:
        if not isinstance(raw, dict):
            continue
        record_id = clean_text(raw.get("record_id"), 80)
        aliases = [clean_text(item, 80) for item in raw.get("aliases", []) if clean_text(item, 80)]
        name = clean_text(raw.get("name"), 80)
        try:
            longitude = float(raw.get("longitude"))
            latitude = float(raw.get("latitude"))
        except (TypeError, ValueError):
            continue
        level = clean_text(raw.get("level"), 20).lower()
        if not record_id or record_id in seen_ids or not name or level not in {"province", "city", "district"}:
            continue
        if not (-180 <= longitude <= 180 and -90 <= latitude <= 90):
            continue
        seen_ids.add(record_id)
        records.append(
            {
                "record_id": record_id,
                "admin_code": clean_text(raw.get("admin_code"), 40),
                "level": level,
                "name": name,
                "province": clean_text(raw.get("province"), 80),
                "city": clean_text(raw.get("city"), 80),
                "district": clean_text(raw.get("district"), 80),
                "aliases": sorted(set([name, *aliases]), key=len, reverse=True),
                "longitude": longitude,
                "latitude": latitude,
                "coordinate_system": clean_text(raw.get("coordinate_system"), 40) or "WGS84",
                "source_version": clean_text(raw.get("source_version"), 80)
                or clean_text(payload.get("source_version"), 80)
                or SOURCE_VERSION,
            }
        )
    if not records:
        raise ProjectSpatialError("地点字典没有可用记录")
    source_title = clean_text(payload.get("source_title"), 180)
    source_authority = clean_text(payload.get("source_authority"), 180)
    license_text = clean_text(payload.get("license"), 240)
    approval_number = clean_text(payload.get("map_approval_number"), 80)
    coordinate_system = clean_text(payload.get("coordinate_system"), 40) or "WGS84"
    requested_formal_map = payload.get("approved_for_formal_map", False) is True
    formal_map_ready = requested_formal_map and all(
        value and "待配置" not in value and "未配置" not in value
        for value in (source_title, source_authority, license_text, approval_number)
    )
    return {
        "source_version": clean_text(payload.get("source_version"), 80) or SOURCE_VERSION,
        "map_provider": clean_text(payload.get("map_provider"), 40).casefold() or MAP_PROVIDER,
        "map_runtime_mode": clean_text(payload.get("map_runtime_mode"), 40).casefold() or MAP_RUNTIME_MODE,
        "source_title": source_title,
        "source_authority": source_authority,
        "license": license_text,
        "map_approval_number": approval_number,
        "coordinate_system": coordinate_system,
        "approved_for_formal_map": formal_map_ready,
        "records": records,
    }


def map_source_metadata(dictionary: dict[str, Any]) -> dict[str, Any]:
    return {
        "source_version": clean_text(dictionary.get("source_version"), 80),
        "map_provider": clean_text(dictionary.get("map_provider"), 40).casefold() or MAP_PROVIDER,
        "map_runtime_mode": clean_text(dictionary.get("map_runtime_mode"), 40).casefold() or MAP_RUNTIME_MODE,
        "source_title": clean_text(dictionary.get("source_title"), 180),
        "source_authority": clean_text(dictionary.get("source_authority"), 180),
        "license": clean_text(dictionary.get("license"), 240),
        "map_approval_number": clean_text(dictionary.get("map_approval_number"), 80),
        "coordinate_system": clean_text(dictionary.get("coordinate_system"), 40),
        "approved_for_formal_map": bool(dictionary.get("approved_for_formal_map", False)),
    }


def _record_alias_matches(text: str, record: dict[str, Any]) -> tuple[int, str] | None:
    normalized = normalize_text(text)
    best: tuple[int, str] | None = None
    for alias in record.get("aliases", []):
        alias_normalized = normalize_text(alias)
        if not alias_normalized:
            continue
        position = normalized.find(alias_normalized)
        if position < 0:
            continue
        candidate = (position, alias)
        if best is None or len(alias_normalized) > len(normalize_text(best[1])):
            best = candidate
    return best


def _specific_candidates(candidates: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Drop an ancestor province when a more specific place in that province is present."""

    selected: list[dict[str, Any]] = []
    for candidate in candidates:
        if candidate.get("level") == "province" and any(
            other.get("province") == candidate.get("name")
            and other.get("level") != "province"
            for other in candidates
        ):
            continue
        selected.append(candidate)
    return selected


def candidates_for_text(text: object, dictionary: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    source = dictionary or load_admin_dictionary()
    matches: list[tuple[int, int, dict[str, Any]]] = []
    for record in source["records"]:
        match = _record_alias_matches(str(text or ""), record)
        if match:
            matches.append((match[0], -len(normalize_text(match[1])), record))
    matches.sort(key=lambda item: (item[0], item[1], str(item[2].get("record_id"))))
    return _specific_candidates([item[2] for item in matches])


def _candidate_public(record: dict[str, Any], *, source_text: str = "", source: dict[str, str] | None = None) -> dict[str, Any]:
    source = source or {}
    return {
        "candidate_id": clean_text(record.get("record_id"), 80),
        "admin_code": clean_text(record.get("admin_code"), 40),
        "level": clean_text(record.get("level"), 20),
        "name": clean_text(record.get("name"), 80),
        "province": clean_text(record.get("province"), 80),
        "city": clean_text(record.get("city"), 80),
        "district": clean_text(record.get("district"), 80),
        "longitude": float(record["longitude"]),
        "latitude": float(record["latitude"]),
        "coordinate_system": clean_text(record.get("coordinate_system"), 40),
        "source_version": clean_text(record.get("source_version"), 80),
        "source_text": clean_text(source_text),
        "source_sheet": clean_text(source.get("sheet_name"), 120),
        "source_cell": clean_text(source.get("cell_address"), 30),
    }


def _field_aliases() -> dict[str, set[str]]:
    return {
        "project_name": {
            normalize_text(item)
            for item in ("项目名称", "工程名称", "项目名", "工程项目", "建设项目名称", "任务名称")
        },
        "location": {
            normalize_text(item)
            for item in (
                "项目地点",
                "项目位置",
                "所在地",
                "地区",
                "省份",
                "行政区划",
                "工程所在地区",
                "建设地点",
                "项目所在地",
                "工程地点",
                "建设区域",
                "所在省市",
            )
        },
        "start": {
            normalize_text(item)
            for item in ("起点", "起始点", "起始位置", "起点位置", "线路起点", "始发地")
        },
        "end": {
            normalize_text(item)
            for item in ("终点", "终止点", "终点位置", "线路终点", "到达地")
        },
        "via": {
            normalize_text(item)
            for item in ("途经", "沿线", "途经地区", "途经地", "沿线地区", "中间点", "路由")
        },
    }


def _merged_grid(sheet: Any) -> dict[tuple[int, int], tuple[Any, str]]:
    merged: dict[tuple[int, int], tuple[Any, str]] = {}
    max_row = min(int(getattr(sheet, "max_row", 0) or 0), MAX_WORKBOOK_ROWS)
    max_column = min(int(getattr(sheet, "max_column", 0) or 0), MAX_WORKBOOK_COLUMNS)
    for merged_range in getattr(sheet, "merged_cells", []).ranges if hasattr(getattr(sheet, "merged_cells", None), "ranges") else []:
        min_row = max(1, int(merged_range.min_row))
        min_col = max(1, int(merged_range.min_col))
        max_row_range = min(max_row, int(merged_range.max_row))
        max_col_range = min(max_column, int(merged_range.max_col))
        if min_row > max_row or min_col > max_column or max_row_range < min_row or max_col_range < min_col:
            continue
        top_value = sheet.cell(min_row, min_col).value
        for row in range(min_row, max_row_range + 1):
            for col in range(min_col, max_col_range + 1):
                if min_col == max_col_range:
                    merged[(row, col)] = (top_value, f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col_range)}{max_row_range}")
                elif row == min_row and col == min_col:
                    merged[(row, col)] = (top_value, f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col_range)}{max_row_range}")
                else:
                    merged[(row, col)] = (None, f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col_range)}{max_row_range}")
    return merged


def _read_workbook_cells(path: Path, selected_sheet: str | None = None) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(path, read_only=False, data_only=True)
    except (OSError, ValueError, KeyError) as exc:
        raise ProjectSpatialValidationError(f"无法读取地点识别输入表：{exc}") from exc
    cells: list[dict[str, Any]] = []
    try:
        sheets = [workbook[selected_sheet]] if selected_sheet and selected_sheet in workbook.sheetnames else list(workbook.worksheets[:MAX_WORKBOOK_SHEETS])
        if selected_sheet and selected_sheet not in workbook.sheetnames:
            raise ProjectSpatialValidationError(f"输入表不存在 sheet：{selected_sheet}")
        for sheet in sheets:
            merged = _merged_grid(sheet)
            max_row = min(int(getattr(sheet, "max_row", 0) or 0), MAX_WORKBOOK_ROWS)
            max_column = min(int(getattr(sheet, "max_column", 0) or 0), MAX_WORKBOOK_COLUMNS)
            for row in range(1, max_row + 1):
                if bool(getattr(sheet.row_dimensions[row], "hidden", False)):
                    continue
                for column in range(1, max_column + 1):
                    value, merged_range = merged.get((row, column), (sheet.cell(row, column).value, ""))
                    if value is None or not clean_text(value):
                        continue
                    cells.append(
                        {
                            "sheet_name": sheet.title,
                            "row": row,
                            "column": column,
                            "cell_address": f"{get_column_letter(column)}{row}",
                            "value": clean_text(value),
                            "merged_range": merged_range,
                        }
                    )
    finally:
        workbook.close()
    return cells


def _cell_index(cells: Iterable[dict[str, Any]]) -> dict[tuple[str, int, int], dict[str, Any]]:
    return {(item["sheet_name"], item["row"], item["column"]): item for item in cells}


def _evidence(
    item: dict[str, Any],
    field_name: str,
    *,
    raw_value: str | None = None,
    method: str,
    source_kind: str | None = None,
) -> dict[str, Any]:
    resolved_source_kind = source_kind or ("excel" if item.get("sheet_name") and item.get("cell_address") else "filename")
    return {
        "field_name": clean_text(field_name, 40),
        "raw_value": clean_text(raw_value if raw_value is not None else item.get("value")),
        "source_text": clean_text(raw_value if raw_value is not None else item.get("value")),
        "sheet_name": clean_text(item.get("sheet_name"), 120),
        "cell_address": clean_text(item.get("cell_address"), 30),
        "merged_range": clean_text(item.get("merged_range"), 40),
        "method": clean_text(method, 40),
        "source_kind": clean_text(resolved_source_kind, 20),
    }


def _adjacent_values(item: dict[str, Any], index: dict[tuple[str, int, int], dict[str, Any]]) -> list[dict[str, Any]]:
    sheet = item["sheet_name"]
    row = int(item["row"])
    column = int(item["column"])
    results: list[dict[str, Any]] = []
    # Horizontal key/value layouts are the common case.  Stop at the first
    # non-empty cell so a neighbouring field cannot become its value.
    for offset in (1, 2, 3):
        candidate = index.get((sheet, row, column + offset))
        if candidate:
            results.append(candidate)
            break
    # Vertical key/value layouts and two-row table headers are both supported.
    for offset in (1, 2):
        candidate = index.get((sheet, row + offset, column))
        if candidate:
            results.append(candidate)
            break
    return results


def _inline_value(value: str) -> str:
    parts = re.split(r"[：:]", value, maxsplit=1)
    return clean_text(parts[1]) if len(parts) == 2 and parts[1].strip() else ""


def _is_ignored_recognition_value(value: object) -> bool:
    normalized = normalize_text(value)
    return normalized in {normalize_text(item) for item in IGNORED_RECOGNITION_VALUES}


def _project_filename_stem(filename: str) -> str:
    stem = Path(filename or "").stem
    stem = re.sub(r"(?:【输出】|【临时】|控制价计算表|控制价报告|[-_ ]v?\d+(?:\.\d+){1,2}).*$", "", stem)
    return clean_text(stem or filename, 160)


def _contains_corridor_marker(text: str) -> bool:
    return bool(re.search(r"(?:→|->|[-—–]|至|到|沿线|途经|干线|管道|联络线|输气线|线路|某标段)", text or ""))


def _place_key(record: dict[str, Any]) -> str:
    return clean_text(record.get("record_id"), 80)


def _build_suggestion(
    *,
    project_name: str,
    filename: str,
    evidence: list[dict[str, Any]],
    location_texts: list[tuple[str, dict[str, Any], bool]],
    dictionary: dict[str, Any],
) -> dict[str, Any]:
    ordered_records: list[dict[str, Any]] = []
    seen: set[str] = set()
    source_candidate_counts: list[tuple[str, list[dict[str, Any]], bool]] = []
    for raw_text, source, explicit in location_texts:
        candidates = candidates_for_text(raw_text, dictionary)
        source_candidate_counts.append((raw_text, candidates, explicit))
        for record in candidates:
            key = _place_key(record)
            if key not in seen:
                seen.add(key)
                ordered_records.append(record)
    inferred_sources: list[dict[str, Any]] = []
    inference_texts = [(project_name, "项目名称"), (_project_filename_stem(filename), "项目名称/文件名")]
    for raw_text, inference_field in inference_texts:
        if not raw_text or (inference_field == "项目名称/文件名" and location_texts):
            continue
        if any(normalize_text(raw_text) == normalize_text(existing[0]) for existing in location_texts):
            continue
        candidates = candidates_for_text(raw_text, dictionary)
        if candidates:
            source = {
                "sheet_name": "",
                "cell_address": "",
                "merged_range": "",
            }
            source_candidate_counts.append((raw_text, candidates, False))
            inferred_sources.append(
                _evidence(
                    source,
                    inference_field,
                    raw_value=raw_text,
                    method="offline_dictionary",
                    source_kind="filename",
                )
            )
            for record in candidates:
                key = _place_key(record)
                if key not in seen:
                    seen.add(key)
                    ordered_records.append(record)
    if inferred_sources:
        evidence = [*evidence, *inferred_sources][:MAX_EVIDENCE_ITEMS]

    conflict = False
    corridor = False
    route_field_names = {
        normalize_text(item)
        for item in ("起点", "起始点", "起始位置", "起点位置", "线路起点", "始发地", "终点", "终止点", "终点位置", "线路终点", "到达地", "途经", "沿线", "途经地区", "途经地", "沿线地区", "中间点", "路由")
    }
    has_route_field = any(
        clean_text(source.get("field_key"), 20) in {"start", "end", "via"}
        or normalize_text(source.get("field_name")) in route_field_names
        for _, source, _ in location_texts
    )
    for raw_text, candidates, explicit in source_candidate_counts:
        if len(candidates) < 2:
            continue
        if _contains_corridor_marker(raw_text) or has_route_field:
            corridor = True
        else:
            conflict = True
    if len(ordered_records) >= 2 and (
        any(_contains_corridor_marker(text) for text, _, _ in source_candidate_counts) or has_route_field
    ):
        corridor = True
    if len(ordered_records) >= 2 and not corridor:
        conflict = True

    if conflict:
        status = "conflict"
        spatial_type = ""
        display_location = "多个地点候选，待人工确认"
        reason = "离线地点字典命中多个非连续地点，不能自动确定项目位置。"
        confidence = 0.35
    elif not ordered_records:
        status = "unlocated"
        spatial_type = ""
        display_location = "未识别地点"
        reason = "未从项目名称、文件名或显式地点字段命中离线地点字典。"
        confidence = 0.0
    elif corridor and len(ordered_records) >= 2:
        status = "proposed"
        spatial_type = "CORRIDOR"
        display_location = " → ".join(clean_text(item.get("name"), 80) for item in ordered_records)
        reason = "命中两个及以上地点，且存在起止点、沿线或管道线路语义；需人工确认路线端点。"
        confidence = 0.86 if any(explicit for _, _, explicit in source_candidate_counts) else 0.76
    elif len(ordered_records) == 1 and ordered_records[0].get("level") == "province":
        status = "proposed"
        spatial_type = "AREA"
        display_location = clean_text(ordered_records[0].get("name"), 80)
        reason = "命中省级行政区，按区域中心点展示；不代表管线边界。"
        confidence = 0.9 if any(explicit for _, _, explicit in source_candidate_counts) else 0.8
    else:
        status = "proposed"
        spatial_type = "POINT"
        display_location = clean_text(ordered_records[0].get("name"), 80)
        reason = "命中市/区县级地点，按代表点展示；需人工确认业务位置。"
        confidence = 0.92 if any(explicit for _, _, explicit in source_candidate_counts) else 0.82

    places: list[dict[str, Any]] = []
    for record in ordered_records:
        source_text = next(
            (
                text
                for text, candidates, _ in source_candidate_counts
                if any(_place_key(item) == _place_key(record) for item in candidates)
            ),
            "",
        )
        source_item = next(
            (
                item
                for item in evidence
                if normalize_text(item.get("source_text")) == normalize_text(source_text)
            ),
            {},
        )
        places.append(_candidate_public(record, source_text=source_text, source=source_item))
    for sequence_no, place in enumerate(places, start=1):
        place["sequence_no"] = sequence_no
    return {
        "available": bool(ordered_records),
        "status": status,
        "recognition_status": status,
        "spatial_type": spatial_type,
        "recognition_method": "excel_explicit" if any(explicit for _, _, explicit in source_candidate_counts) else "offline_dictionary",
        "display_location": display_location,
        "confidence": round(confidence, 3),
        "reason": reason,
        "project_name": clean_text(project_name, 160),
        "source_filename": clean_text(filename, 180),
        "source_version": clean_text(dictionary.get("source_version"), 80),
        "places": places,
        "evidence": evidence[:MAX_EVIDENCE_ITEMS],
        "requires_confirmation": status != "confirmed",
        "map_source": map_source_metadata(dictionary),
    }


def recognize_project_identity(
    *,
    project_name: str = "",
    source_filename: str = "",
    explicit_values: list[dict[str, Any]] | None = None,
    dictionary: dict[str, Any] | None = None,
) -> dict[str, Any]:
    source = dictionary or load_admin_dictionary()
    evidence: list[dict[str, Any]] = []
    location_texts: list[tuple[str, dict[str, Any], bool]] = []
    for item in explicit_values or []:
        if not isinstance(item, dict):
            continue
        raw_value = clean_text(item.get("raw_value") or item.get("source_text"))
        if not raw_value:
            continue
        field_name = clean_text(item.get("field_name"), 40) or "项目地点"
        item_evidence = {
            "field_name": field_name,
            "raw_value": raw_value,
            "source_text": raw_value,
            "sheet_name": clean_text(item.get("sheet_name"), 120),
            "cell_address": clean_text(item.get("cell_address"), 30),
            "merged_range": clean_text(item.get("merged_range"), 40),
            "method": clean_text(item.get("method"), 40) or "excel_explicit",
            "source_kind": clean_text(item.get("source_kind"), 20) or "excel",
        }
        evidence.append(item_evidence)
        field_key = clean_text(item.get("field_key"), 20)
        item_evidence["field_key"] = field_key
        field_name_key = normalize_text(field_name)
        location_aliases = set().union(*(_field_aliases()[key] for key in ("location", "start", "end", "via")))
        known_field = field_name_key in location_aliases
        if field_key in {"location", "start", "end", "via"} or known_field:
            location_texts.append((raw_value, item_evidence, True))
    return _build_suggestion(
        project_name=clean_text(project_name, 160),
        filename=clean_text(source_filename, 180),
        evidence=evidence,
        location_texts=location_texts,
        dictionary=source,
    )


def recognize_workbook(
    path: Path,
    *,
    project_name_hint: str = "",
    source_filename: str = "",
    sheet_name: str | None = None,
    dictionary: dict[str, Any] | None = None,
) -> dict[str, Any]:
    source = dictionary or load_admin_dictionary()
    filename = source_filename or Path(path).name
    cells = _read_workbook_cells(Path(path), selected_sheet=sheet_name)
    index = _cell_index(cells)
    aliases = _field_aliases()
    project_name = clean_text(project_name_hint, 160)
    explicit_values: list[dict[str, Any]] = []
    explicit_name_evidence: list[dict[str, Any]] = []
    for item in cells:
        normalized = normalize_text(item["value"])
        label = re.split(r"[：:]", item["value"], maxsplit=1)[0].strip()
        normalized_label = normalize_text(label)
        matched_key = next(
            (key for key, values in aliases.items() if normalized in values or normalized_label in values),
            "",
        )
        inline = _inline_value(item["value"])
        if matched_key == "project_name":
            if inline and not project_name and not _is_ignored_recognition_value(inline):
                project_name = inline
                explicit_name_evidence.append(_evidence(item, "项目名称", raw_value=inline, method="excel_explicit"))
            for adjacent in _adjacent_values(item, index):
                adjacent_value = clean_text(adjacent.get("value"), 160)
                if not project_name and not _is_ignored_recognition_value(adjacent_value):
                    project_name = adjacent_value
                    explicit_name_evidence.append(_evidence(adjacent, "项目名称", method="excel_explicit"))
                break
        if matched_key in {"location", "start", "end", "via"}:
            values = [inline] if inline else []
            if not values:
                values = [clean_text(adjacent.get("value")) for adjacent in _adjacent_values(item, index)]
            for value in values[:2]:
                if _is_ignored_recognition_value(value) or normalize_text(value) in aliases.get(matched_key, set()):
                    continue
                explicit_values.append(
                    {
                        "field_key": matched_key,
                        "field_name": label,
                        "raw_value": value,
                        "source_text": value,
                        "sheet_name": item["sheet_name"],
                        "cell_address": item["cell_address"],
                        "merged_range": item.get("merged_range", ""),
                        "method": "excel_explicit",
                        "source_kind": "excel",
                    }
                )
    if not project_name and not explicit_values:
        project_name = _project_filename_stem(filename)
    suggestion = recognize_project_identity(
        project_name=project_name,
        source_filename=filename,
        explicit_values=explicit_values,
        dictionary=source,
    )
    if explicit_name_evidence:
        suggestion["evidence"] = [*explicit_name_evidence, *suggestion.get("evidence", [])][:MAX_EVIDENCE_ITEMS]
    suggestion["input_sheet"] = clean_text(sheet_name or "", 120)
    suggestion["inspected_cell_count"] = len(cells)
    return suggestion


def _validate_spatial_type(value: object) -> str:
    spatial_type = clean_text(value, 20).upper()
    if spatial_type not in SPATIAL_TYPES:
        raise ProjectSpatialValidationError("空间类型必须是 POINT、AREA 或 CORRIDOR")
    return spatial_type


def _records_by_id(dictionary: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {clean_text(item.get("record_id"), 80): item for item in dictionary.get("records", [])}


def _places_from_payload(
    payload: object,
    *,
    dictionary: dict[str, Any],
    spatial_type: str,
) -> list[dict[str, Any]]:
    if not isinstance(payload, list):
        raise ProjectSpatialValidationError("地点列表必须是数组")
    if not 1 <= len(payload) <= 20:
        raise ProjectSpatialValidationError("地点列表数量不符合要求")
    records = _records_by_id(dictionary)
    places: list[dict[str, Any]] = []
    seen: set[str] = set()
    for sequence_no, raw in enumerate(payload, start=1):
        if not isinstance(raw, dict):
            raise ProjectSpatialValidationError("地点候选必须是对象")
        if any(key in raw for key in ("longitude", "latitude", "经度", "纬度")):
            raise ProjectSpatialValidationError("坐标必须由服务器从受信地点字典读取")
        candidate_id = clean_text(raw.get("candidate_id") or raw.get("record_id") or raw.get("admin_code"), 80)
        record = records.get(candidate_id)
        if record is None:
            record = next((item for item in records.values() if clean_text(item.get("admin_code"), 40) == candidate_id), None)
        if record is None:
            raise ProjectSpatialValidationError(f"地点候选不在受信字典中：{candidate_id}")
        key = _place_key(record)
        if key in seen:
            raise ProjectSpatialValidationError("地点候选不能重复")
        seen.add(key)
        places.append(
            {
                "sequence_no": sequence_no,
                "record": record,
                "source_text": clean_text(raw.get("source_text")),
                "source_sheet": clean_text(raw.get("source_sheet") or raw.get("sheet_name"), 120),
                "source_cell": clean_text(raw.get("source_cell") or raw.get("cell_address"), 30),
            }
        )
    if spatial_type in {"POINT", "AREA"} and len(places) != 1:
        raise ProjectSpatialValidationError(f"{spatial_type} 必须恰好有一个地点")
    if spatial_type == "CORRIDOR" and len(places) < 2:
        raise ProjectSpatialValidationError("CORRIDOR 至少需要两个地点")
    return places


def _display_location(spatial_type: str, places: list[dict[str, Any]]) -> str:
    if spatial_type == "CORRIDOR":
        return " → ".join(clean_text(item["record"].get("name"), 80) for item in places)
    return clean_text(places[0]["record"].get("name"), 80)


SCHEMA = """
CREATE TABLE IF NOT EXISTS project_spatial_profiles (
    project_id TEXT PRIMARY KEY,
    spatial_type TEXT NOT NULL DEFAULT '',
    recognition_status TEXT NOT NULL DEFAULT 'unlocated',
    recognition_method TEXT NOT NULL DEFAULT '',
    display_location TEXT NOT NULL DEFAULT '',
    evidence_json TEXT NOT NULL DEFAULT '{}',
    source_version TEXT NOT NULL DEFAULT '',
    revision INTEGER NOT NULL DEFAULT 0,
    confirmed_by TEXT NOT NULL DEFAULT '',
    confirmed_at TEXT NOT NULL DEFAULT '',
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL,
    FOREIGN KEY(project_id) REFERENCES projects(project_id) ON DELETE CASCADE
);
CREATE TABLE IF NOT EXISTS project_spatial_places (
    place_id TEXT PRIMARY KEY,
    project_id TEXT NOT NULL,
    sequence_no INTEGER NOT NULL,
    admin_code TEXT NOT NULL,
    province TEXT NOT NULL DEFAULT '',
    city TEXT NOT NULL DEFAULT '',
    district TEXT NOT NULL DEFAULT '',
    longitude REAL NOT NULL,
    latitude REAL NOT NULL,
    coordinate_system TEXT NOT NULL,
    source_text TEXT NOT NULL DEFAULT '',
    source_sheet TEXT NOT NULL DEFAULT '',
    source_cell TEXT NOT NULL DEFAULT '',
    FOREIGN KEY(project_id) REFERENCES project_spatial_profiles(project_id) ON DELETE CASCADE,
    UNIQUE(project_id, sequence_no)
);
CREATE TABLE IF NOT EXISTS project_spatial_revisions (
    revision_id TEXT PRIMARY KEY,
    project_id TEXT NOT NULL,
    action TEXT NOT NULL,
    before_json TEXT NOT NULL DEFAULT '{}',
    after_json TEXT NOT NULL DEFAULT '{}',
    actor TEXT NOT NULL DEFAULT '',
    reason TEXT NOT NULL DEFAULT '',
    idempotency_key TEXT NOT NULL DEFAULT '',
    created_at TEXT NOT NULL,
    FOREIGN KEY(project_id) REFERENCES projects(project_id) ON DELETE CASCADE
);
CREATE TABLE IF NOT EXISTS project_spatial_pending_candidates (
    pending_id TEXT PRIMARY KEY,
    project_id TEXT NOT NULL,
    fingerprint TEXT NOT NULL,
    candidate_json TEXT NOT NULL DEFAULT '{}',
    source TEXT NOT NULL DEFAULT '',
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL,
    consumed_at TEXT NOT NULL DEFAULT '',
    FOREIGN KEY(project_id) REFERENCES projects(project_id) ON DELETE CASCADE,
    UNIQUE(project_id, fingerprint)
);
"""


SPATIAL_MIGRATION_COLUMNS: dict[str, dict[str, str]] = {
    "project_spatial_profiles": {
        "spatial_type": "TEXT NOT NULL DEFAULT ''",
        "recognition_status": "TEXT NOT NULL DEFAULT 'unlocated'",
        "recognition_method": "TEXT NOT NULL DEFAULT ''",
        "display_location": "TEXT NOT NULL DEFAULT ''",
        "evidence_json": "TEXT NOT NULL DEFAULT '{}'",
        "source_version": "TEXT NOT NULL DEFAULT ''",
        "revision": "INTEGER NOT NULL DEFAULT 0",
        "confirmed_by": "TEXT NOT NULL DEFAULT ''",
        "confirmed_at": "TEXT NOT NULL DEFAULT ''",
        "created_at": "TEXT NOT NULL DEFAULT ''",
        "updated_at": "TEXT NOT NULL DEFAULT ''",
    },
    "project_spatial_places": {
        "project_id": "TEXT NOT NULL DEFAULT ''",
        "sequence_no": "INTEGER NOT NULL DEFAULT 1",
        "admin_code": "TEXT NOT NULL DEFAULT ''",
        "province": "TEXT NOT NULL DEFAULT ''",
        "city": "TEXT NOT NULL DEFAULT ''",
        "district": "TEXT NOT NULL DEFAULT ''",
        "longitude": "REAL NOT NULL DEFAULT 0",
        "latitude": "REAL NOT NULL DEFAULT 0",
        "coordinate_system": "TEXT NOT NULL DEFAULT 'WGS84'",
        "source_text": "TEXT NOT NULL DEFAULT ''",
        "source_sheet": "TEXT NOT NULL DEFAULT ''",
        "source_cell": "TEXT NOT NULL DEFAULT ''",
    },
    "project_spatial_revisions": {
        "project_id": "TEXT NOT NULL DEFAULT ''",
        "action": "TEXT NOT NULL DEFAULT ''",
        "before_json": "TEXT NOT NULL DEFAULT '{}'",
        "after_json": "TEXT NOT NULL DEFAULT '{}'",
        "actor": "TEXT NOT NULL DEFAULT ''",
        "reason": "TEXT NOT NULL DEFAULT ''",
        "idempotency_key": "TEXT NOT NULL DEFAULT ''",
        "created_at": "TEXT NOT NULL DEFAULT ''",
    },
    "project_spatial_pending_candidates": {
        "project_id": "TEXT NOT NULL DEFAULT ''",
        "fingerprint": "TEXT NOT NULL DEFAULT ''",
        "candidate_json": "TEXT NOT NULL DEFAULT '{}'",
        "source": "TEXT NOT NULL DEFAULT ''",
        "created_at": "TEXT NOT NULL DEFAULT ''",
        "updated_at": "TEXT NOT NULL DEFAULT ''",
        "consumed_at": "TEXT NOT NULL DEFAULT ''",
    },
}


DEMO_PROJECTS: list[dict[str, Any]] = [
    {
        "project_id": "DEMO-XJ-01",
        "project_name": "新疆北部天然气管道勘察项目",
        "spatial_type": "AREA",
        "display_location": "新疆维吾尔自治区",
        "candidate_ids": ["adm-xinjiang"],
        "risk_level": "low",
        "status": "completed",
    },
    {
        "project_id": "DEMO-CQSC-02",
        "project_name": "川渝天然气干线某标段",
        "spatial_type": "CORRIDOR",
        "display_location": "成都 → 重庆",
        "candidate_ids": ["adm-chengdu", "adm-chongqing"],
        "risk_level": "medium",
        "status": "pending_review",
    },
    {
        "project_id": "DEMO-SD-03",
        "project_name": "山东潍坊输气管道测量项目",
        "spatial_type": "POINT",
        "display_location": "潍坊市",
        "candidate_ids": ["adm-weifang"],
        "risk_level": "low",
        "status": "completed",
    },
    {
        "project_id": "DEMO-HE-04",
        "project_name": "河北廊坊管线改线勘察项目",
        "spatial_type": "POINT",
        "display_location": "廊坊市",
        "candidate_ids": ["adm-langfang"],
        "risk_level": "high",
        "status": "pending_review",
    },
    {
        "project_id": "DEMO-JS-05",
        "project_name": "江苏沿江输气管道工程",
        "spatial_type": "CORRIDOR",
        "display_location": "南京 → 苏州",
        "candidate_ids": ["adm-nanjing", "adm-suzhou"],
        "risk_level": "medium",
        "status": "completed",
    },
    {
        "project_id": "DEMO-GD-06",
        "project_name": "广东珠三角管网勘察项目",
        "spatial_type": "AREA",
        "display_location": "广东省",
        "candidate_ids": ["adm-guangdong"],
        "risk_level": "low",
        "status": "completed",
    },
    {
        "project_id": "DEMO-SX-07",
        "project_name": "陕北干线某段复测项目",
        "spatial_type": "CORRIDOR",
        "display_location": "西安 → 榆林",
        "candidate_ids": ["adm-xian", "adm-yulin"],
        "risk_level": "high",
        "status": "pending_review",
    },
    {
        "project_id": "DEMO-YN-08",
        "project_name": "云贵联络线勘察项目",
        "spatial_type": "CORRIDOR",
        "display_location": "昆明 → 贵阳",
        "candidate_ids": ["adm-kunming", "adm-guiyang"],
        "risk_level": "medium",
        "status": "completed",
    },
]


def demo_map_items(dictionary: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    source = dictionary or load_admin_dictionary()
    records = _records_by_id(source)
    items: list[dict[str, Any]] = []
    for project in DEMO_PROJECTS:
        places = []
        for sequence_no, candidate_id in enumerate(project["candidate_ids"], start=1):
            record = records.get(candidate_id)
            if not record:
                continue
            place = _candidate_public(record, source_text="演示样例", source={})
            place["sequence_no"] = sequence_no
            places.append(place)
        items.append(
            {
                "project_id": project["project_id"],
                "project_name": project["project_name"],
                "spatial_type": project["spatial_type"],
                "recognition_status": "confirmed",
                "recognition_method": "demo_fixture",
                "display_location": project["display_location"],
                "status": project["status"],
                "status_label": {"completed": "已完成", "pending_review": "待复核"}.get(project["status"], project["status"]),
                "risk_level": project["risk_level"],
                "risk_label": {"high": "高风险", "medium": "待关注", "low": "低风险"}.get(project["risk_level"], "未知"),
                "is_demo": True,
                "map_ready": map_runtime_enabled(source),
                "map_source": map_source_metadata(source),
                "places": places,
                "place_count": len(places),
                "latest_run_id": "",
                "updated_at": "2026-09-12T00:00:00+08:00",
            }
        )
    return items


class ProjectSpatialStore:
    """Spatial tables in the existing project-ledger database."""

    def __init__(self, db_path: Path, dictionary_path: Path | None = None) -> None:
        self.db_path = Path(db_path)
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self.dictionary = load_admin_dictionary(dictionary_path)
        self._init_db()

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.db_path, timeout=15)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys=ON")
        connection.execute("PRAGMA busy_timeout=15000")
        return connection

    def _init_db(self) -> None:
        with self._connect() as connection:
            connection.executescript(SCHEMA)
            self._migrate_schema(connection)

    @staticmethod
    def _migrate_schema(connection: sqlite3.Connection) -> None:
        """Add spatial columns/indexes when an older runtime DB is reused."""

        for table_name, columns in SPATIAL_MIGRATION_COLUMNS.items():
            existing = {
                str(row[1])
                for row in connection.execute(f"PRAGMA table_info({table_name})").fetchall()
            }
            for column_name, definition in columns.items():
                if column_name not in existing:
                    connection.execute(
                        f"ALTER TABLE {table_name} ADD COLUMN {column_name} {definition}"
                    )
        # A malformed/duplicated legacy audit key must not prevent the new
        # idempotency index from being recreated. Keep the earliest audit row
        # for each non-empty key and leave empty keys unrestricted.
        connection.execute(
            """
            DELETE FROM project_spatial_revisions
            WHERE idempotency_key<>''
              AND rowid NOT IN (
                  SELECT MIN(rowid)
                  FROM project_spatial_revisions
                  WHERE idempotency_key<>''
                  GROUP BY project_id,idempotency_key
              )
            """
        )
        connection.execute(
            """
            CREATE UNIQUE INDEX IF NOT EXISTS idx_project_spatial_revisions_idempotency
            ON project_spatial_revisions(project_id, idempotency_key)
            WHERE idempotency_key<>''
            """
        )
        connection.execute(
            """
            CREATE INDEX IF NOT EXISTS idx_project_spatial_places_project_id
            ON project_spatial_places(project_id, sequence_no)
            """
        )
        connection.execute(
            """
            CREATE INDEX IF NOT EXISTS idx_project_spatial_revisions_project_id
            ON project_spatial_revisions(project_id, created_at DESC)
            """
        )
        connection.execute(
            """
            CREATE INDEX IF NOT EXISTS idx_project_spatial_pending_project_id
            ON project_spatial_pending_candidates(project_id, updated_at DESC)
            """
        )

    def _project_exists(self, project_id: str, connection: sqlite3.Connection) -> str:
        try:
            clean_project_id = validate_project_id(project_id)
        except ProjectNotFoundError as exc:
            raise ProjectSpatialNotFound("项目编号无效") from exc
        row = connection.execute(
            "SELECT project_id FROM projects WHERE project_id=? AND archived=0", (clean_project_id,)
        ).fetchone()
        if not row:
            raise ProjectSpatialNotFound("项目不存在或已归档")
        return clean_project_id

    def _stored_place_payload(self, place: sqlite3.Row) -> dict[str, Any] | None:
        trusted_record = next(
            (
                item
                for item in self.dictionary.get("records", [])
                if clean_text(item.get("admin_code"), 40) == clean_text(place["admin_code"], 40)
            ),
            None,
        )
        if trusted_record is None:
            return None
        public = _candidate_public(
            trusted_record,
            source_text=str(place["source_text"] or ""),
            source={"sheet_name": place["source_sheet"], "cell_address": place["source_cell"]},
        )
        return {
            "place_id": str(place["place_id"]),
            "sequence_no": int(place["sequence_no"]),
            **public,
            "source_text": str(place["source_text"] or ""),
            "source_sheet": str(place["source_sheet"] or ""),
            "source_cell": str(place["source_cell"] or ""),
        }

    @staticmethod
    def _pending_payload(row: sqlite3.Row) -> dict[str, Any] | None:
        try:
            candidate = json.loads(str(row["candidate_json"] or "{}"))
        except json.JSONDecodeError:
            return None
        if not isinstance(candidate, dict):
            return None
        return {
            "pending_id": str(row["pending_id"]),
            "source": str(row["source"] or ""),
            "created_at": str(row["created_at"] or ""),
            "updated_at": str(row["updated_at"] or ""),
            **candidate,
        }

    def _profile_payload(self, connection: sqlite3.Connection, project_id: str) -> dict[str, Any] | None:
        row = connection.execute(
            "SELECT * FROM project_spatial_profiles WHERE project_id=?", (project_id,)
        ).fetchone()
        if not row:
            return None
        places = connection.execute(
            "SELECT * FROM project_spatial_places WHERE project_id=? ORDER BY sequence_no ASC", (project_id,)
        ).fetchall()
        stored_places = [
            stored_place
            for place in places
            if (stored_place := self._stored_place_payload(place)) is not None
        ]
        pending_rows = connection.execute(
            """
            SELECT * FROM project_spatial_pending_candidates
            WHERE project_id=? AND consumed_at=''
            ORDER BY updated_at DESC,pending_id ASC
            LIMIT 20
            """,
            (project_id,),
        ).fetchall()
        pending_candidates = [
            pending for row in pending_rows if (pending := self._pending_payload(row)) is not None
        ]
        evidence: object
        try:
            evidence = json.loads(str(row["evidence_json"] or "{}"))
        except json.JSONDecodeError:
            evidence = {}
        return {
            "project_id": project_id,
            "spatial_type": str(row["spatial_type"] or ""),
            "recognition_status": str(row["recognition_status"] or "unlocated"),
            "recognition_method": str(row["recognition_method"] or ""),
            "display_location": str(row["display_location"] or ""),
            "evidence": evidence,
            "source_version": str(row["source_version"] or ""),
            "revision": int(row["revision"] or 0),
            "confirmed_by": str(row["confirmed_by"] or ""),
            "confirmed_at": str(row["confirmed_at"] or ""),
            "created_at": str(row["created_at"] or ""),
            "updated_at": str(row["updated_at"] or ""),
            "places": stored_places,
            "pending_candidates": pending_candidates,
            "map_source": map_source_metadata(self.dictionary),
            "map_ready": bool(
                str(row["recognition_status"]) == "confirmed"
                and stored_places
                and str(row["source_version"] or "") == map_source_metadata(self.dictionary)["source_version"]
                and all(place["coordinate_system"] for place in stored_places)
                and map_runtime_enabled(self.dictionary)
            ),
        }

    def get(self, project_id: str) -> dict[str, Any]:
        with self._connect() as connection:
            clean_project_id = self._project_exists(project_id, connection)
            payload = self._profile_payload(connection, clean_project_id)
        if payload is None:
            return {
                "project_id": clean_project_id,
                "spatial_type": "",
                "recognition_status": "unlocated",
                "recognition_method": "",
                "display_location": "未识别地点",
                "evidence": {},
                "source_version": map_source_metadata(self.dictionary)["source_version"],
                "revision": 0,
                "confirmed_by": "",
                "confirmed_at": "",
                "created_at": "",
                "updated_at": "",
                "places": [],
                "pending_candidates": [],
                "map_source": map_source_metadata(self.dictionary),
                "map_ready": False,
            }
        return payload

    def _revision_payload(self, connection: sqlite3.Connection, project_id: str) -> dict[str, Any]:
        return self._profile_payload(connection, project_id) or {
            "project_id": project_id,
            "recognition_status": "unlocated",
            "places": [],
            "pending_candidates": [],
        }

    def _candidate_snapshot(
        self,
        *,
        spatial_type: str,
        places: list[dict[str, Any]],
        method: str,
        display_location: str,
        evidence: object,
        reason: object,
        revision: int,
    ) -> dict[str, Any]:
        public_places = []
        for place in places:
            record = place["record"]
            public_place = _candidate_public(
                record,
                source_text=place.get("source_text", ""),
                source={
                    "sheet_name": place.get("source_sheet", ""),
                    "cell_address": place.get("source_cell", ""),
                },
            )
            public_place["sequence_no"] = int(place["sequence_no"])
            public_places.append(public_place)
        return {
            "spatial_type": spatial_type,
            "recognition_status": "proposed",
            "recognition_method": method,
            "display_location": display_location,
            "evidence": evidence,
            "reason": clean_text(reason, 240),
            "source_version": map_source_metadata(self.dictionary)["source_version"],
            "revision": revision,
            "places": public_places,
            "requires_confirmation": True,
            "map_source": map_source_metadata(self.dictionary),
        }

    def validate_analysis_context(self, project_id: str, payload: object) -> dict[str, Any]:
        """Validate that a spatial analysis request is tied to current state."""

        if not isinstance(payload, dict):
            raise ProjectSpatialValidationError("地点分析请求必须是对象")
        with self._connect() as connection:
            clean_project_id = self._project_exists(project_id, connection)
            profile = self._profile_payload(connection, clean_project_id)
            current_revision = int((profile or {}).get("revision") or 0)
            if "revision" in payload and payload.get("revision") is not None:
                try:
                    requested_revision = int(payload.get("revision"))
                except (TypeError, ValueError) as exc:
                    raise ProjectSpatialValidationError("地点分析 revision 必须是整数") from exc
                if requested_revision != current_revision:
                    raise ProjectSpatialConflict("地点分析基于旧版本地点资料，请刷新后重试")
            run_id = clean_text(payload.get("run_id"), 120)
            if run_id and not re.fullmatch(r"[A-Za-z0-9_.:-]{1,120}", run_id):
                raise ProjectSpatialValidationError("地点分析 run_id 格式无效")
            project_row = connection.execute(
                "SELECT latest_run_id FROM projects WHERE project_id=? AND archived=0",
                (clean_project_id,),
            ).fetchone()
            latest_run_id = clean_text(project_row["latest_run_id"] if project_row else "", 120)
            if run_id and run_id != latest_run_id:
                raise ProjectSpatialConflict("地点分析只能使用项目最新一版 Run")
            return {
                "project_id": clean_project_id,
                "revision": current_revision,
                "run_id": latest_run_id,
            }

    def propose(
        self,
        project_id: str,
        payload: dict[str, Any],
        *,
        actor: str = "system",
        idempotency_key: str = "",
    ) -> dict[str, Any]:
        spatial_type = _validate_spatial_type(payload.get("spatial_type"))
        if any(key in payload for key in ("longitude", "latitude", "经度", "纬度")):
            raise ProjectSpatialValidationError("坐标必须由服务器从受信地点字典读取")
        current_source_version = map_source_metadata(self.dictionary)["source_version"]
        supplied_source_version = clean_text(payload.get("source_version"), 80)
        if supplied_source_version and supplied_source_version != current_source_version:
            raise ProjectSpatialValidationError("地点候选来源版本已过期，请重新读取候选")
        places = _places_from_payload(payload.get("places"), dictionary=self.dictionary, spatial_type=spatial_type)
        clean_actor = clean_text(actor, 80) or "system"
        clean_idempotency = clean_text(idempotency_key or payload.get("idempotency_key"), 120)
        evidence = _safe_json(payload.get("evidence") or {}, limit=MAX_TEXT_LENGTH)
        method = clean_text(payload.get("recognition_method"), 40) or "manual"
        display_location = clean_text(payload.get("display_location"), 160) or _display_location(spatial_type, places)
        with self._connect() as connection:
            connection.execute("BEGIN IMMEDIATE")
            clean_project_id = self._project_exists(project_id, connection)
            if clean_idempotency:
                existing = connection.execute(
                    "SELECT 1 FROM project_spatial_revisions WHERE project_id=? AND idempotency_key=?",
                    (clean_project_id, clean_idempotency),
                ).fetchone()
                if existing:
                    return self._revision_payload(connection, clean_project_id)
            before = self._revision_payload(connection, clean_project_id)
            current_revision = int(before.get("revision") or 0)
            timestamp = now_iso()
            if str(before.get("recognition_status") or "") == "confirmed":
                candidate = self._candidate_snapshot(
                    spatial_type=spatial_type,
                    places=places,
                    method=method,
                    display_location=display_location,
                    evidence=evidence,
                    reason=payload.get("reason"),
                    revision=current_revision,
                )
                fingerprint = hashlib.sha256(
                    json.dumps(
                        {
                            "spatial_type": spatial_type,
                            "candidate_ids": [item["candidate_id"] for item in candidate["places"]],
                            "source_version": candidate["source_version"],
                        },
                        ensure_ascii=False,
                        sort_keys=True,
                    ).encode("utf-8")
                ).hexdigest()
                current_candidate_ids = [
                    clean_text(item.get("admin_code"), 40)
                    for item in before.get("places", [])
                    if isinstance(item, dict)
                ]
                discovered_candidate_ids = [
                    clean_text(item.get("admin_code"), 40)
                    for item in candidate["places"]
                    if isinstance(item, dict)
                ]
                if spatial_type == clean_text(before.get("spatial_type"), 20).upper() and current_candidate_ids == discovered_candidate_ids:
                    return {
                        **before,
                        "spatial_update": {
                            "status": "confirmed_unchanged",
                            "message": "新文件地点与当前已确认地点一致，未创建重复候选。",
                        },
                    }
                consumed_pending = connection.execute(
                    """
                    SELECT pending_id FROM project_spatial_pending_candidates
                    WHERE project_id=? AND fingerprint=? AND consumed_at<>''
                    LIMIT 1
                    """,
                    (clean_project_id, fingerprint),
                ).fetchone()
                if consumed_pending:
                    return {
                        **before,
                        "spatial_update": {
                            "status": "confirmed_unchanged",
                            "message": "新候选与当前已确认地点一致，未重复创建待复核记录。",
                        },
                    }
                connection.execute(
                    """
                    INSERT INTO project_spatial_pending_candidates(
                        pending_id,project_id,fingerprint,candidate_json,source,created_at,updated_at,consumed_at
                    ) VALUES(?,?,?,?,?,?,?, '')
                    ON CONFLICT(project_id,fingerprint) DO UPDATE SET updated_at=excluded.updated_at
                    WHERE project_spatial_pending_candidates.consumed_at=''
                    """,
                    (
                        f"spp_{uuid4().hex[:24]}",
                        clean_project_id,
                        fingerprint,
                        json.dumps(candidate, ensure_ascii=False),
                        clean_actor,
                        timestamp,
                        timestamp,
                    ),
                )
                after = self._revision_payload(connection, clean_project_id)
                connection.execute(
                    """
                    INSERT INTO project_spatial_revisions(
                        revision_id,project_id,action,before_json,after_json,actor,reason,idempotency_key,created_at
                    ) VALUES(?,?,?,?,?,?,?,?,?)
                    """,
                    (
                        f"spr_{uuid4().hex[:24]}",
                        clean_project_id,
                        "discover_candidate",
                        json.dumps(before, ensure_ascii=False),
                        json.dumps({**after, "discovered_candidate": candidate}, ensure_ascii=False),
                        clean_actor,
                        clean_text(payload.get("reason"), 240),
                        clean_idempotency,
                        timestamp,
                    ),
                )
                return {
                    **after,
                    "spatial_update": {
                        "status": "new_candidate",
                        "message": "已有人工确认地点，已记录新候选，未覆盖当前确认结果。",
                        "pending_candidate": candidate,
                    },
                }
            connection.execute(
                """
                INSERT INTO project_spatial_profiles(
                    project_id,spatial_type,recognition_status,recognition_method,display_location,
                    evidence_json,source_version,revision,confirmed_by,confirmed_at,created_at,updated_at
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)
                ON CONFLICT(project_id) DO UPDATE SET
                    spatial_type=excluded.spatial_type,
                    recognition_status='proposed',
                    recognition_method=excluded.recognition_method,
                    display_location=excluded.display_location,
                    evidence_json=excluded.evidence_json,
                    source_version=excluded.source_version,
                    revision=excluded.revision,
                    confirmed_by='',
                    confirmed_at='',
                    updated_at=excluded.updated_at
                """,
                (
                    clean_project_id,
                    spatial_type,
                    "proposed",
                    method,
                    display_location,
                    json.dumps(evidence, ensure_ascii=False),
                    map_source_metadata(self.dictionary)["source_version"],
                    current_revision + 1,
                    "",
                    "",
                    timestamp,
                    timestamp,
                ),
            )
            connection.execute("DELETE FROM project_spatial_places WHERE project_id=?", (clean_project_id,))
            for place in places:
                record = place["record"]
                connection.execute(
                    """
                    INSERT INTO project_spatial_places(
                        place_id,project_id,sequence_no,admin_code,province,city,district,
                        longitude,latitude,coordinate_system,source_text,source_sheet,source_cell
                    ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)
                    """,
                    (
                        f"spl_{uuid4().hex[:24]}",
                        clean_project_id,
                        int(place["sequence_no"]),
                        clean_text(record.get("admin_code"), 40),
                        clean_text(record.get("province"), 80),
                        clean_text(record.get("city"), 80),
                        clean_text(record.get("district"), 80),
                        float(record["longitude"]),
                        float(record["latitude"]),
                        clean_text(record.get("coordinate_system"), 40),
                        place["source_text"],
                        place["source_sheet"],
                        place["source_cell"],
                    ),
                )
            after = self._revision_payload(connection, clean_project_id)
            connection.execute(
                """
                INSERT INTO project_spatial_revisions(
                    revision_id,project_id,action,before_json,after_json,actor,reason,idempotency_key,created_at
                ) VALUES(?,?,?,?,?,?,?,?,?)
                """,
                (
                    f"spr_{uuid4().hex[:24]}",
                    clean_project_id,
                    "propose",
                    json.dumps(before, ensure_ascii=False),
                    json.dumps(after, ensure_ascii=False),
                    clean_actor,
                    clean_text(payload.get("reason"), 240),
                    clean_idempotency,
                    timestamp,
                ),
            )
            return after

    def update(
        self,
        project_id: str,
        payload: dict[str, Any],
        *,
        actor: str = "human",
        idempotency_key: str = "",
    ) -> dict[str, Any]:
        action = clean_text(payload.get("action"), 20).lower()
        if action not in {"confirm", "correct", "reject", "clear"}:
            raise ProjectSpatialValidationError("空间更新 action 必须是 confirm、correct、reject 或 clear")
        clean_actor = clean_text(actor or payload.get("actor"), 80) or "human"
        clean_idempotency = clean_text(idempotency_key or payload.get("idempotency_key"), 120)
        with self._connect() as connection:
            connection.execute("BEGIN IMMEDIATE")
            clean_project_id = self._project_exists(project_id, connection)
            if clean_idempotency:
                existing = connection.execute(
                    "SELECT 1 FROM project_spatial_revisions WHERE project_id=? AND idempotency_key=?",
                    (clean_project_id, clean_idempotency),
                ).fetchone()
                if existing:
                    return self._revision_payload(connection, clean_project_id)
            before = self._revision_payload(connection, clean_project_id)
            expected_revision = payload.get("expected_revision")
            if expected_revision is not None:
                try:
                    expected_revision_number = int(expected_revision)
                except (TypeError, ValueError) as exc:
                    raise ProjectSpatialValidationError("expected_revision 必须是整数") from exc
                if expected_revision_number != int(before.get("revision") or 0):
                    raise ProjectSpatialConflict("地点资料已被其他操作更新，请刷新后再确认")
            timestamp = now_iso()
            profile_exists = connection.execute(
                "SELECT 1 FROM project_spatial_profiles WHERE project_id=?", (clean_project_id,)
            ).fetchone() is not None
            if action == "clear":
                connection.execute("DELETE FROM project_spatial_profiles WHERE project_id=?", (clean_project_id,))
                connection.execute(
                    "UPDATE project_spatial_pending_candidates SET consumed_at=?,updated_at=? WHERE project_id=? AND consumed_at=''",
                    (timestamp, timestamp, clean_project_id),
                )
            else:
                current_type = clean_text(before.get("spatial_type"), 20).upper()
                current_places = before.get("places") if isinstance(before.get("places"), list) else []
                if action == "correct":
                    current_type = _validate_spatial_type(payload.get("spatial_type"))
                    normalized_places = _places_from_payload(payload.get("places"), dictionary=self.dictionary, spatial_type=current_type)
                    supplied_source_version = clean_text(payload.get("source_version"), 80)
                    current_source_version = map_source_metadata(self.dictionary)["source_version"]
                    if supplied_source_version and supplied_source_version != current_source_version:
                        raise ProjectSpatialValidationError("地点候选来源版本已过期，请重新读取候选")
                    current_places = normalized_places
                    evidence = _safe_json(payload.get("evidence") or before.get("evidence") or {}, limit=MAX_TEXT_LENGTH)
                    display_location = clean_text(payload.get("display_location"), 160) or _display_location(current_type, normalized_places)
                    source_version = map_source_metadata(self.dictionary)["source_version"]
                    if profile_exists:
                        connection.execute(
                            """
                            UPDATE project_spatial_profiles
                            SET spatial_type=?,recognition_status='proposed',recognition_method='manual',display_location=?,
                                evidence_json=?,source_version=?,revision=?,confirmed_by='',confirmed_at='',updated_at=?
                            WHERE project_id=?
                            """,
                            (
                                current_type,
                                display_location,
                                json.dumps(evidence, ensure_ascii=False),
                                source_version,
                                int(before.get("revision") or 0) + 1,
                                timestamp,
                                clean_project_id,
                            ),
                        )
                    else:
                        connection.execute(
                            """
                            INSERT INTO project_spatial_profiles(
                                project_id,spatial_type,recognition_status,recognition_method,display_location,
                                evidence_json,source_version,revision,confirmed_by,confirmed_at,created_at,updated_at
                            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)
                            """,
                            (
                                clean_project_id,
                                current_type,
                                "proposed",
                                "manual",
                                display_location,
                                json.dumps(evidence, ensure_ascii=False),
                                source_version,
                                1,
                                "",
                                "",
                                timestamp,
                                timestamp,
                            ),
                        )
                    connection.execute("DELETE FROM project_spatial_places WHERE project_id=?", (clean_project_id,))
                    for place in normalized_places:
                        record = place["record"]
                        connection.execute(
                            """
                            INSERT INTO project_spatial_places(
                                place_id,project_id,sequence_no,admin_code,province,city,district,
                                longitude,latitude,coordinate_system,source_text,source_sheet,source_cell
                            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)
                            """,
                            (
                                f"spl_{uuid4().hex[:24]}", clean_project_id, int(place["sequence_no"]),
                                clean_text(record.get("admin_code"), 40), clean_text(record.get("province"), 80),
                                clean_text(record.get("city"), 80), clean_text(record.get("district"), 80),
                                float(record["longitude"]), float(record["latitude"]),
                                clean_text(record.get("coordinate_system"), 40), place["source_text"],
                                place["source_sheet"], place["source_cell"],
                            ),
                        )
                    connection.execute(
                        "UPDATE project_spatial_pending_candidates SET consumed_at=?,updated_at=? WHERE project_id=? AND consumed_at=''",
                        (timestamp, timestamp, clean_project_id),
                    )
                elif action == "confirm":
                    if current_type not in SPATIAL_TYPES or not current_places:
                        raise ProjectSpatialValidationError("没有可确认的地点候选")
                    required_count = 2 if current_type == "CORRIDOR" else 1
                    if len(current_places) < required_count:
                        raise ProjectSpatialValidationError("地点数量不满足空间类型约束")
                    connection.execute(
                        """
                        UPDATE project_spatial_profiles
                        SET recognition_status='confirmed',recognition_method=CASE WHEN recognition_method='' THEN 'manual' ELSE recognition_method END,
                            revision=?,confirmed_by=?,confirmed_at=?,updated_at=?
                        WHERE project_id=?
                        """,
                        (
                            int(before.get("revision") or 0) + 1,
                            clean_actor,
                            timestamp,
                            timestamp,
                            clean_project_id,
                        ),
                    )
                elif action == "reject":
                    connection.execute(
                        """
                        UPDATE project_spatial_profiles
                        SET recognition_status='rejected',revision=?,confirmed_by='',confirmed_at='',updated_at=?
                        WHERE project_id=?
                        """,
                        (int(before.get("revision") or 0) + 1, timestamp, clean_project_id),
                    )
                    connection.execute(
                        "UPDATE project_spatial_pending_candidates SET consumed_at=?,updated_at=? WHERE project_id=? AND consumed_at=''",
                        (timestamp, timestamp, clean_project_id),
                    )
            after = self._revision_payload(connection, clean_project_id)
            connection.execute(
                """
                INSERT INTO project_spatial_revisions(
                    revision_id,project_id,action,before_json,after_json,actor,reason,idempotency_key,created_at
                ) VALUES(?,?,?,?,?,?,?,?,?)
                """,
                (
                    f"spr_{uuid4().hex[:24]}", clean_project_id, action,
                    json.dumps(before, ensure_ascii=False), json.dumps(after, ensure_ascii=False),
                    clean_actor, clean_text(payload.get("reason"), 240), clean_idempotency, timestamp,
                ),
            )
            return after

    def map_items(
        self,
        *,
        scope: str = "real",
        keyword: str = "",
        spatial_type: str = "",
        status: str = "",
        risk: str = "",
        date_from: str = "",
        date_to: str = "",
        limit: int = 500,
    ) -> dict[str, Any]:
        clean_scope = clean_text(scope, 20).lower() or "real"
        if clean_scope not in {"real", "demo"}:
            raise ProjectSpatialValidationError("地图 scope 只能是 real 或 demo")
        if clean_scope == "demo":
            items = demo_map_items(self.dictionary)
            source = "demo_fixture"
        else:
            items = self._real_map_items()
            source = "project_ledger"
        try:
            start_date = date.fromisoformat(clean_text(date_from, 10)) if clean_text(date_from, 10) else None
            end_date = date.fromisoformat(clean_text(date_to, 10)) if clean_text(date_to, 10) else None
        except ValueError as exc:
            raise ProjectSpatialValidationError("地图日期筛选必须是 YYYY-MM-DD") from exc
        if start_date and end_date and start_date > end_date:
            raise ProjectSpatialValidationError("地图日期筛选的起始日期不能晚于结束日期")
        keyword_norm = normalize_text(keyword)
        requested_type = clean_text(spatial_type, 20).upper()
        if requested_type and requested_type not in {"ALL", *SPATIAL_TYPES}:
            raise ProjectSpatialValidationError("地图空间类型筛选无效")
        requested_status = clean_text(status, 40).lower()
        if requested_status and requested_status not in SPATIAL_STATUSES | {"processing", "pending_review", "completed", "returned", "failed"}:
            raise ProjectSpatialValidationError("地图项目状态筛选无效")
        requested_risk = clean_text(risk, 20).lower()
        if requested_risk and requested_risk not in {"high", "medium", "low", "unknown"}:
            raise ProjectSpatialValidationError("地图风险筛选无效")
        result: list[dict[str, Any]] = []
        for item in items:
            if keyword_norm and keyword_norm not in normalize_text(item.get("project_name")) and keyword_norm not in normalize_text(item.get("display_location")):
                continue
            if requested_type and requested_type not in {"ALL", clean_text(item.get("spatial_type"), 20).upper()}:
                continue
            if requested_status and requested_status != clean_text(item.get("status"), 40).lower() and requested_status != clean_text(item.get("recognition_status"), 40).lower():
                continue
            if requested_risk and requested_risk != clean_text(item.get("risk_level"), 20).lower():
                continue
            item_day = _date_only(item.get("created_at") or item.get("updated_at"))
            if start_date and (item_day is None or item_day < start_date):
                continue
            if end_date and (item_day is None or item_day > end_date):
                continue
            result.append(item)
        max_items = max(1, min(500, int(limit or 500)))
        return {
            "scope": clean_scope,
            "source": source,
            "items": result[:max_items],
            "total": len(result),
            "truncated": len(result) > max_items,
            "map_source": map_source_metadata(self.dictionary),
            "map_ready": map_runtime_enabled(self.dictionary),
            "degraded": not map_runtime_enabled(self.dictionary),
            "message": (
                "测试项目直接使用 AntV L7；默认使用 OpenFreeMap Positron 测试底图，地图异常时降级为示意层。"
                if map_runtime_enabled(self.dictionary)
                else "AntV L7 测试运行模式未启用，当前仅提供地点清单与示意层。"
            ),
        }

    def _real_map_items(self) -> list[dict[str, Any]]:
        with self._connect() as connection:
            rows = connection.execute(
                """
                SELECT p.project_id,p.project_name,p.status,p.latest_run_id,p.updated_at,
                       p.created_at,p.source_type,r.risk_high,r.risk_low,r.review_rows,
                       r.input_rows,r.matched_rows,r.experience_hint_rows,
                       (SELECT COUNT(*) FROM project_artifacts a WHERE a.project_id=p.project_id) AS artifact_count,
                       s.spatial_type,s.recognition_status,
                       s.recognition_method,s.display_location,s.source_version
                FROM project_spatial_profiles s
                JOIN projects p ON p.project_id=s.project_id AND p.archived=0
                LEFT JOIN project_runs r ON r.run_id=p.latest_run_id
                WHERE s.recognition_status='confirmed'
                ORDER BY p.updated_at DESC,p.project_id ASC
                """
            ).fetchall()
            items: list[dict[str, Any]] = []
            for row in rows:
                places = connection.execute(
                    "SELECT * FROM project_spatial_places WHERE project_id=? ORDER BY sequence_no ASC",
                    (row["project_id"],),
                ).fetchall()
                public_places = [
                    stored_place
                    for place in places
                    if (stored_place := self._stored_place_payload(place)) is not None
                ]
                risk_level = "high" if int(row["risk_high"] or 0) > 0 else "low" if int(row["risk_low"] or 0) > 0 else "unknown"
                items.append(
                    {
                        "project_id": str(row["project_id"]),
                        "project_name": str(row["project_name"]),
                        "spatial_type": str(row["spatial_type"]),
                        "recognition_status": str(row["recognition_status"]),
                        "recognition_method": str(row["recognition_method"] or ""),
                        "display_location": str(row["display_location"] or ""),
                        "status": str(row["status"]),
                        "status_label": {"processing": "处理中", "pending_review": "待复核", "completed": "已完成", "returned": "已退回", "failed": "失败"}.get(str(row["status"]), str(row["status"])),
                        "risk_level": risk_level,
                        "risk_label": {"high": "高风险", "low": "低风险", "unknown": "未运行预警"}.get(risk_level, risk_level),
                        "is_demo": False,
                        "map_ready": bool(public_places) and map_runtime_enabled(self.dictionary),
                        "map_source": map_source_metadata(self.dictionary),
                        "places": public_places,
                        "place_count": len(public_places),
                        "latest_run_id": str(row["latest_run_id"] or ""),
                        "created_at": str(row["created_at"] or ""),
                        "updated_at": str(row["updated_at"] or ""),
                        "source_type": str(row["source_type"] or ""),
                        "review_rows": int(row["review_rows"] or 0),
                        "artifact_summary": {
                            "artifact_count": int(row["artifact_count"] or 0),
                            "input_rows": int(row["input_rows"] or 0),
                            "matched_rows": int(row["matched_rows"] or 0),
                            "experience_hint_rows": int(row["experience_hint_rows"] or 0),
                            "review_rows": int(row["review_rows"] or 0),
                        },
                    }
                )
            return items

    def unlocated(
        self,
        *,
        keyword: str = "",
        date_from: str = "",
        date_to: str = "",
        limit: int = 200,
        page: int = 1,
        page_size: int | None = None,
    ) -> dict[str, Any]:
        with self._connect() as connection:
            rows = connection.execute(
                """
                SELECT p.project_id,p.project_name,p.status,p.latest_run_id,p.created_at,p.updated_at,
                       s.recognition_status,s.display_location,s.revision
                FROM projects p
                LEFT JOIN project_spatial_profiles s ON s.project_id=p.project_id
                WHERE p.archived=0 AND (s.project_id IS NULL OR s.recognition_status<>'confirmed')
                ORDER BY p.updated_at DESC,p.project_id ASC
                """
            ).fetchall()
        keyword_norm = normalize_text(keyword)
        try:
            start_date = date.fromisoformat(clean_text(date_from, 10)) if clean_text(date_from, 10) else None
            end_date = date.fromisoformat(clean_text(date_to, 10)) if clean_text(date_to, 10) else None
        except ValueError as exc:
            raise ProjectSpatialValidationError("未定位项目日期筛选必须是 YYYY-MM-DD") from exc
        if start_date and end_date and start_date > end_date:
            raise ProjectSpatialValidationError("未定位项目日期筛选的起始日期不能晚于结束日期")
        items = [
            {
                "project_id": str(row["project_id"]),
                "project_name": str(row["project_name"]),
                "status": str(row["status"]),
                "latest_run_id": str(row["latest_run_id"] or ""),
                "created_at": str(row["created_at"] or ""),
                "updated_at": str(row["updated_at"] or ""),
                "recognition_status": str(row["recognition_status"] or "unlocated"),
                "display_location": str(row["display_location"] or "未识别地点"),
                "revision": int(row["revision"] or 0),
                "needs_review": True,
            }
            for row in rows
            if (
                not keyword_norm
                or keyword_norm in normalize_text(row["project_name"])
                or keyword_norm in normalize_text(row["display_location"] or "")
            )
            and (
                (not start_date and not end_date)
                or (
                    (item_day := _date_only(row["created_at"])) is not None
                    and (not start_date or item_day >= start_date)
                    and (not end_date or item_day <= end_date)
                )
            )
        ]
        try:
            page_number = max(1, int(page or 1))
            max_items = max(1, min(200, int(page_size or limit or 200)))
        except (TypeError, ValueError) as exc:
            raise ProjectSpatialValidationError("未定位项目分页参数无效") from exc
        start = (page_number - 1) * max_items
        end = start + max_items
        return {
            "items": items[start:end],
            "total": len(items),
            "truncated": len(items) > max_items,
            "page": page_number,
            "page_size": max_items,
            "has_more": end < len(items),
            "map_source": map_source_metadata(self.dictionary),
        }

    def candidates(self, *, keyword: str = "", limit: int = 30) -> dict[str, Any]:
        keyword_norm = normalize_text(keyword)
        records = self.dictionary.get("records", [])
        items = []
        for record in records:
            searchable = "".join(
                [
                    clean_text(record.get("name")),
                    clean_text(record.get("province")),
                    clean_text(record.get("city")),
                    clean_text(record.get("district")),
                    "".join(clean_text(alias) for alias in record.get("aliases", [])),
                ]
            )
            if keyword_norm and keyword_norm not in normalize_text(searchable):
                continue
            item = _candidate_public(record)
            item.pop("source_text", None)
            item.pop("source_sheet", None)
            item.pop("source_cell", None)
            items.append(item)
        max_items = max(1, min(100, int(limit or 30)))
        return {
            "items": items[:max_items],
            "total": len(items),
            "truncated": len(items) > max_items,
            "map_source": map_source_metadata(self.dictionary),
        }
