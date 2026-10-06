from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..fill_engine import FillEngine
from ..schemas import FIELD_COLUMNS


REQUIRED_FIELDS = ("要素1", "单位")
MAX_CONTEXT_ROWS = 3


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _header_map(headers: list[Any]) -> dict[str, int]:
    return {
        _text(value): index
        for index, value in enumerate(headers, start=1)
        if _text(value)
    }


def _source_cell(
    source_sheet: Any,
    value_sheet: Any,
    row: int,
    column: int,
    merged_value_map: dict[tuple[int, int], Any],
) -> dict[str, Any]:
    cell = source_sheet.cell(row=row, column=column)
    cached_cell = value_sheet.cell(row=row, column=column)
    value = merged_value_map.get((row, column), cached_cell.value)
    return {
        "sheet": source_sheet.title,
        "cell": f"{get_column_letter(column)}{row}",
        "range": f"{get_column_letter(column)}{row}",
        "value": value,
        "formula": cell.value if isinstance(cell.value, str) and cell.value.startswith("=") else "",
        "cached_value": cached_cell.value,
    }


def _context_rows(value_sheet: Any, row: int, max_column: int) -> list[dict[str, Any]]:
    context: list[dict[str, Any]] = []
    start = max(1, row - MAX_CONTEXT_ROWS)
    for context_row in range(start, row):
        values = [_text(value_sheet.cell(row=context_row, column=column).value) for column in range(1, max_column + 1)]
        compact = [value for value in values if value]
        if compact:
            context.append(
                {
                    "excel_row": context_row,
                    "values": values,
                    "source_cells": {
                        f"{get_column_letter(column)}{context_row}": {
                            "sheet": value_sheet.title,
                            "cell": f"{get_column_letter(column)}{context_row}",
                            "range": f"{get_column_letter(column)}{context_row}",
                            "value": value_sheet.cell(row=context_row, column=column).value,
                        }
                        for column in range(1, max_column + 1)
                        if value_sheet.cell(row=context_row, column=column).value not in (None, "")
                    },
                }
            )
    return context


def read_workbook_structure(
    path: str | Path,
    *,
    column_mapping: dict[str, str] | None = None,
    header_row: int = 1,
    sheet_configs: list[dict[str, Any]] | None = None,
    merge_vertical_cells: bool = True,
    merge_horizontal_cells: bool = True,
    only_match_rows_with_value: bool = False,
    match_value_filter_field: str = "数量",
    max_input_rows: int = 500,
) -> dict[str, Any]:
    max_input_rows = max(1, int(max_input_rows or 500))
    workbook = load_workbook(path, data_only=False)
    value_workbook = load_workbook(path, data_only=True)
    try:
        configs = sheet_configs or [
            {
                "sheet_name": workbook.active.title,
                "enabled": True,
                "header_row": header_row,
                "column_mapping": column_mapping,
                "merge_vertical_cells": merge_vertical_cells,
                "merge_horizontal_cells": merge_horizontal_cells,
                "only_match_rows_with_value": only_match_rows_with_value,
                "match_value_filter_field": match_value_filter_field,
            }
        ]
        sheets: list[dict[str, Any]] = []
        total_business_rows = 0
        total_scanned_rows = 0
        total_excluded_rows = 0
        truncated = False
        missing_fields: set[str] = set()
        standard_mapping_count = 0

        for config in configs:
            if config.get("enabled") is False:
                continue
            sheet_name = _text(config.get("sheet_name")) or workbook.active.title
            if sheet_name not in workbook.sheetnames:
                sheets.append({"sheet_name": sheet_name, "error": "sheet_not_found"})
                continue
            source_sheet = workbook[sheet_name]
            value_sheet = value_workbook[sheet_name]
            current_header_row = int(config.get("header_row") or header_row)
            headers = [_text(cell.value) for cell in value_sheet[current_header_row]]
            headers = headers[: max(1, source_sheet.max_column)]
            header_map = _header_map(headers)
            current_mapping = config.get("column_mapping") or column_mapping
            field_map = FillEngine._resolve_field_map(header_map, current_mapping)
            current_missing = [field for field in REQUIRED_FIELDS if field not in field_map]
            missing_fields.update(current_missing)
            if not current_missing:
                standard_mapping_count += 1
            merged_value_map = FillEngine._build_merged_value_map(
                value_sheet,
                merge_vertical_cells=bool(config.get("merge_vertical_cells", merge_vertical_cells)),
                merge_horizontal_cells=bool(config.get("merge_horizontal_cells", merge_horizontal_cells)),
            )
            filter_column = (
                FillEngine._find_value_filter_column(headers, str(config.get("match_value_filter_field") or match_value_filter_field))
                if bool(config.get("only_match_rows_with_value", only_match_rows_with_value))
                else None
            )
            if bool(config.get("only_match_rows_with_value", only_match_rows_with_value)) and filter_column is None:
                raise ValueError(f"{sheet_name} 未找到指定列：{config.get('match_value_filter_field') or match_value_filter_field}")
            records: list[dict[str, Any]] = []
            excluded: list[dict[str, Any]] = []
            for excel_row in range(current_header_row + 1, source_sheet.max_row + 1):
                total_scanned_rows += 1
                if len(records) >= max_input_rows:
                    truncated = True
                    break
                values = {
                    field: FillEngine._read_mapped_value(
                        value_sheet,
                        excel_row,
                        field_map.get(field, 0),
                        merged_value_map,
                    )
                    for field in FIELD_COLUMNS
                }
                raw_values = [_text(value_sheet.cell(row=excel_row, column=column).value) for column in range(1, source_sheet.max_column + 1)]
                row_needs_semantic_mapping = "要素1" not in field_map
                row_is_blank = not any(raw_values)
                row_is_total = any(value.replace(" ", "").startswith(("合计", "小计", "总计")) for value in raw_values if value)
                header_signature = [value for value in headers if value]
                row_signature = [value for value in raw_values[: len(headers)] if value]
                row_is_repeated_header = bool(header_signature) and row_signature == header_signature
                if row_is_repeated_header:
                    excluded.append({"excel_row": excel_row, "reason": "repeated_header_row"})
                    total_excluded_rows += 1
                    continue
                if FillEngine._is_ignored_row(values.get("要素1")) and not (
                    row_needs_semantic_mapping and not row_is_blank and not row_is_total
                ):
                    excluded.append({"excel_row": excel_row, "reason": "blank_or_total_row"})
                    total_excluded_rows += 1
                    continue
                if filter_column is not None:
                    filter_value = FillEngine._read_mapped_value(value_sheet, excel_row, filter_column, merged_value_map)
                    if not FillEngine._has_value_for_matching_filter(filter_value):
                        excluded.append({"excel_row": excel_row, "reason": "filter_field_empty_or_zero"})
                        total_excluded_rows += 1
                        continue
                source_cells = {
                    field: _source_cell(
                        source_sheet,
                        value_sheet,
                        excel_row,
                        field_map[field],
                        merged_value_map,
                    )
                    for field in FIELD_COLUMNS
                    if field_map.get(field, 0)
                }
                raw_source_cells = {
                    f"{get_column_letter(column)}{excel_row}": _source_cell(
                        source_sheet,
                        value_sheet,
                        excel_row,
                        column,
                        merged_value_map,
                    )
                    for column in range(1, source_sheet.max_column + 1)
                    if (
                        value_sheet.cell(row=excel_row, column=column).value not in (None, "")
                        or source_sheet.cell(row=excel_row, column=column).value not in (None, "")
                    )
                }
                records.append(
                    {
                        "record_id": f"{sheet_name}!{excel_row}",
                        "sheet_name": sheet_name,
                        "excel_row": excel_row,
                        "values": values,
                        "raw_values": raw_values,
                        "raw_text": " | ".join(value for value in raw_values if value),
                        "source_cells": source_cells,
                        "raw_source_cells": raw_source_cells,
                        "context_rows": _context_rows(value_sheet, excel_row, min(source_sheet.max_column, 12)),
                        "candidate_fields": {},
                    }
                )
            total_business_rows += len(records)
            sheets.append(
                {
                    "sheet_name": sheet_name,
                    "header_row": current_header_row,
                    "headers": headers,
                    "field_map": field_map,
                    "missing_fields": current_missing,
                    "records": records,
                    "excluded_rows": excluded[:100],
                    "coverage": {
                        "scanned_rows": source_sheet.max_row - current_header_row,
                        "business_rows": len(records),
                        "excluded_rows": len(excluded),
                        "truncated": bool(len(records) >= max_input_rows and source_sheet.max_row > current_header_row),
                    },
                }
            )

        return {
            "schema_version": "workbook-reader-v1",
            "input_sha256": hashlib.sha256(Path(path).read_bytes()).hexdigest(),
            "structure_kind": "standard" if not missing_fields and standard_mapping_count else "needs_semantic_mapping",
            "missing_fields": sorted(missing_fields),
            "sheets": sheets,
            "records": [record for sheet in sheets for record in sheet.get("records", [])],
            "coverage": {
                "business_rows": total_business_rows,
                "scanned_rows": total_scanned_rows,
                "excluded_rows": total_excluded_rows,
                "truncated": truncated,
                "sheet_count": len(sheets),
            },
        }
    finally:
        workbook.close()
        value_workbook.close()
