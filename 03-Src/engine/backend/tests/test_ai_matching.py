from __future__ import annotations

import asyncio
from pathlib import Path
import json
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

import app.ai_matching.orchestrator as orchestrator
from app.main import app
from app.ai_matching.client import AiMatchingClient, AiMatchingClientConfig
from app.ai_matching.contracts import empty_ai_state
from app.ai_matching.orchestrator import analyze_run, confirm_proposals, prepare_run
from app.ai_matching.orchestrator import _validate_model_output
from app.ai_matching.store import load_ai_state, save_ai_state
from app.ai_matching.workbook_reader import read_workbook_structure
from app.llm import LlmConfig
from app.professional_skills import SkillRuntimeContext


def _write_structurally_ambiguous_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "费用明细"
    sheet.append(["项目内容", "工作类别", "计量单位", "价格", "数量"])
    sheet.append(["控制测量", "GPS 测量", "点", 3000, 2])
    sheet.append(["", "", "", "", ""])
    workbook.save(path)
    workbook.close()


def _write_multirow_semantic_workbook(path: Path, row_count: int = 3) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "费用明细"
    sheet.append(["项目内容", "计量单位", "数量"])
    for index in range(row_count):
        sheet.append([f"控制测量-{index + 1}", "点", index + 1])
    workbook.save(path)
    workbook.close()


def _write_candidate_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "表2"
    sheet.append(["项目内容（要素1）", "专业类别（要素2）", "比例尺（要素4）", "复杂程度（要素5）", "计量单位（单位）", "数量", "价格"])
    sheet.append(["岩土工程勘察", "地质测绘", "比例-1:500", "简单", "km2", 1, None])
    workbook.save(path)
    workbook.close()


def _minimal_context(tmp_path: Path) -> SkillRuntimeContext:
    return SkillRuntimeContext(
        skill_id="test-skill",
        skill_version="1.0.0",
        manifest_hash="test-manifest",
        input_profile={"extensions": [".xlsx"]},
        processor_id="survey-measurement-v1",
        knowledge_base_path=tmp_path / "missing-kb.xlsx",
        rule_assets={
            "technicalRules": (tmp_path / "technical.xlsx",),
            "physicalRules": (tmp_path / "physical.xlsx",),
        },
        risk_profile={"experiencePool": tmp_path / "missing-pool.xlsx"},
        knowledge_sources=(),
        report_template_path=tmp_path / "missing-template.docx",
        validation_profile={},
        capabilities={"pricing": True},
    )


def test_reader_keeps_nonblank_rows_when_semantic_mapping_is_needed(tmp_path):
    workbook_path = tmp_path / "ambiguous.xlsx"
    _write_structurally_ambiguous_workbook(workbook_path)

    structure = read_workbook_structure(
        workbook_path,
        header_row=1,
        only_match_rows_with_value=True,
        match_value_filter_field="数量",
    )

    assert structure["structure_kind"] == "needs_semantic_mapping"
    assert structure["coverage"]["business_rows"] == 1
    assert structure["records"][0]["record_id"] == "费用明细!2"
    assert structure["records"][0]["source_cells"] == {}


def test_reader_excludes_repeated_header_rows_from_business_records(tmp_path):
    workbook_path = tmp_path / "repeated-header.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "费用明细"
    headers = ["项目内容", "计量单位", "数量"]
    sheet.append(headers)
    sheet.append(["控制测量", "点", 1])
    sheet.append(headers)
    sheet.append(["GPS 测量", "点", 2])
    workbook.save(workbook_path)
    workbook.close()

    structure = read_workbook_structure(workbook_path, header_row=1, only_match_rows_with_value=False)

    assert [record["excel_row"] for record in structure["records"]] == [2, 4]
    assert structure["coverage"]["excluded_rows"] == 1
    assert structure["sheets"][0]["excluded_rows"] == [{"excel_row": 3, "reason": "repeated_header_row"}]


def test_mechanical_prepare_is_offline_and_does_not_call_model(tmp_path, monkeypatch):
    workbook_path = tmp_path / "input.xlsx"
    _write_structurally_ambiguous_workbook(workbook_path)
    context = _minimal_context(tmp_path)

    def fail_if_called(*args, **kwargs):
        raise AssertionError("机械模式不应触碰模型客户端")

    monkeypatch.setattr("app.ai_matching.client.AiMatchingClient.call_json", fail_if_called)
    state = prepare_run(
        tmp_path / "job",
        workbook_path,
        requested_mode="mechanical",
        options={"ai_matching": {"defaultMode": "ai_enhanced"}},
        context=context,
    )

    assert state["requested_mode"] == "mechanical"
    assert state["effective_mode"] == "mechanical"
    assert state["ai_stage"] == "not_requested"
    assert load_ai_state(tmp_path / "job")["ai_stage"] == "not_requested"


def test_ai_prepare_creates_audit_sidecar_before_analysis(tmp_path):
    workbook_path = tmp_path / "input.xlsx"
    _write_structurally_ambiguous_workbook(workbook_path)
    job_dir = tmp_path / "new-job"

    state = prepare_run(
        job_dir,
        workbook_path,
        requested_mode="ai_enhanced",
        options={"ai_matching": {"maxInputRows": 25}},
        context=_minimal_context(tmp_path),
    )

    assert state["ai_stage"] == "awaiting_analysis"
    audit_path = job_dir / "ai-matching-audit.jsonl"
    assert audit_path.exists()
    assert json.loads(audit_path.read_text(encoding="utf-8").splitlines()[0])["event"] == "run_prepared"


def test_ai_client_freezes_bounded_runtime_settings_without_exposing_secret():
    client = AiMatchingClient.from_options(
        {
            "ai_matching": {
                "requestTimeoutSeconds": 999,
                "totalModelBudgetSeconds": 0,
                "maxTransientRetries": 99,
                "maxOutputTokens": 99999,
            }
        }
    )

    public = client.public_config()
    assert public["request_timeout_seconds"] == 60.0
    assert public["total_budget_seconds"] == 1.0
    assert public["max_transient_retries"] == 2
    assert public["max_tokens"] == 4000
    assert "api_key" not in public
    assert "DEEPSEEK_API_KEY" not in str(public)


def test_ai_client_requests_non_thinking_json_output(monkeypatch):
    captured = {}

    class _Response:
        def __enter__(self):
            return self

        def __exit__(self, *args):
            return None

        def read(self):
            return b'{"choices":[{"message":{"content":"{}"}}]}'

    def fake_urlopen(request, timeout):
        captured["payload"] = json.loads(request.data.decode("utf-8"))
        captured["timeout"] = timeout
        return _Response()

    monkeypatch.setattr("app.ai_matching.client.urlopen", fake_urlopen)
    monkeypatch.setattr("app.ai_matching.client.LlmConfig.resolved_api_key", lambda self: "test-key")
    client = AiMatchingClient(AiMatchingClientConfig(max_tokens=800))
    raw = client._request(LlmConfig(model=client.config.model, max_tokens=800), [], 12)

    assert raw == "{}"
    assert captured["payload"]["thinking"] == {"type": "disabled"}
    assert captured["payload"]["response_format"] == {"type": "json_object"}
    assert captured["payload"]["max_tokens"] == 800


def test_ai_client_async_request_is_cancellable(monkeypatch):
    class _HangingAsyncClient:
        def __init__(self, *args, **kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            return None

        async def post(self, *args, **kwargs):
            await asyncio.sleep(60)

    monkeypatch.setattr("app.ai_matching.client.httpx.AsyncClient", _HangingAsyncClient)
    monkeypatch.setattr("app.ai_matching.client.LlmConfig.resolved_api_key", lambda self: "test-key")

    async def run_and_cancel():
        client = AiMatchingClient(
            AiMatchingClientConfig(
                request_timeout_seconds=20,
                total_budget_seconds=60,
                max_transient_retries=0,
                max_tokens=800,
            )
        )
        task = asyncio.create_task(client.call_json([], started_at=time.monotonic()))
        await asyncio.sleep(0)
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task

    asyncio.run(run_and_cancel())


def test_ai_analysis_uses_bounded_blocks_and_preserves_order(tmp_path, monkeypatch):
    workbook_path = tmp_path / "multirow.xlsx"
    _write_multirow_semantic_workbook(workbook_path, row_count=3)
    job_dir = tmp_path / "job"
    options = {
        "ai_matching": {
            "maxRowsPerBlock": 1,
            "maxConcurrency": 2,
        }
    }
    context = _minimal_context(tmp_path)
    prepare_run(job_dir, workbook_path, requested_mode="ai_enhanced", options=options, context=context)
    active = 0
    max_active = 0
    calls = 0

    async def fake_call_json(self, messages, *, started_at):
        nonlocal active, max_active, calls
        active += 1
        max_active = max(max_active, active)
        calls += 1
        self.calls += 1
        await asyncio.sleep(0.01)
        active -= 1
        return {"mapping": {"要素1": "项目内容", "单位": "计量单位"}, "proposals": []}

    monkeypatch.setattr(orchestrator.AiMatchingClient, "call_json", fake_call_json)
    result = asyncio.run(analyze_run(job_dir, workbook_path, options=options, context=context))

    assert result["ai_stage"] == "awaiting_confirmation"
    assert result["metrics"]["block_count"] == 3
    assert result["metrics"]["successful_blocks"] == 3
    assert result["metrics"]["failed_blocks"] == 0
    assert result["metrics"]["model_calls"] == 3
    assert calls == 3
    assert max_active <= 2
    assert result["model_result"]["mapping"] == {"要素1": "项目内容", "单位": "计量单位"}


def test_ai_analysis_reuses_validated_cache_without_calling_model_again(tmp_path, monkeypatch):
    workbook_path = tmp_path / "cache-input.xlsx"
    _write_structurally_ambiguous_workbook(workbook_path)
    job_dir = tmp_path / "job"
    options = {"ai_matching": {"maxRowsPerBlock": 25}}
    context = _minimal_context(tmp_path)
    prepare_run(job_dir, workbook_path, requested_mode="ai_enhanced", options=options, context=context)
    calls = 0

    async def fake_call_json(self, messages, *, started_at):
        nonlocal calls
        calls += 1
        self.calls += 1
        return {"mapping": {"要素1": "项目内容", "单位": "计量单位"}, "proposals": []}

    monkeypatch.setattr(orchestrator.AiMatchingClient, "call_json", fake_call_json)
    first = asyncio.run(analyze_run(job_dir, workbook_path, options=options, context=context))
    second = asyncio.run(analyze_run(job_dir, workbook_path, options=options, context=context))

    assert first["model_result"]["source"] == "model_live"
    assert second["model_result"]["source"] == "model_cache"
    assert second["metrics"]["cache_hits"] == 1
    assert second["metrics"]["model_calls"] == 0
    assert calls == 1


def test_confirmation_requires_known_proposal_and_records_revision(tmp_path):
    job_dir = tmp_path / "job"
    state = empty_ai_state(requested_mode="ai_enhanced", effective_mode="ai_enhanced")
    state.update(
        {
            "ai_stage": "awaiting_confirmation",
            "operation_revision": 3,
            "proposals": [
                {
                    "proposal_id": "aip_known",
                    "record_id": "费用明细!2",
                    "field": "price",
                    "candidate_id": "kb-2",
                    "status": "ready_for_confirmation",
                    "decision": "needs_review",
                    "evidence_cells": ["费用明细!A2"],
                    "missing_conditions": [],
                    "candidate_snapshot": {"id": "kb-2", "field": "price", "value": 100},
                    "can_confirm": True,
                }
            ],
        }
    )
    save_ai_state(job_dir, state)

    with pytest.raises(ValueError, match="未知 AI 提案"):
        confirm_proposals(job_dir, proposal_ids=["not-from-server"], expected_revision=3)

    confirmed = confirm_proposals(
        job_dir,
        proposal_ids=["aip_known"],
        mapping={"要素1": "项目内容"},
        expected_revision=3,
    )

    assert confirmed["ai_stage"] == "confirmed"
    assert confirmed["confirmed_proposal_ids"] == ["aip_known"]
    assert confirmed["confirmed_mapping"] == {"要素1": "项目内容"}
    assert confirmed["operation_revision"] == 4


def test_ai_output_normalizes_string_missing_condition_without_character_split():
    structure = {
        "records": [{"record_id": "费用明细!2", "candidate_fields": {}}],
        "sheets": [{"headers": ["项目内容", "计量单位"]}],
    }
    _, proposals = _validate_model_output(
        {
            "mapping": {},
            "proposals": [
                {
                    "record_id": "费用明细!2",
                    "field": "单位",
                    "candidate_id": None,
                    "missing_conditions": "单位字段缺失",
                    "decision": "needs_review",
                }
            ],
        },
        structure,
        [],
    )

    assert proposals[0]["missing_conditions"] == ["单位字段缺失"]


def test_ai_process_keeps_preview_pending_until_mapping_confirmation(tmp_path, monkeypatch):
    workbook_path = tmp_path / "ambiguous-input.xlsx"
    _write_structurally_ambiguous_workbook(workbook_path)

    async def fake_call_json(self, messages, *, started_at):
        self.calls += 1
        return {"mapping": {"要素1": "项目内容", "单位": "计量单位"}, "proposals": []}

    monkeypatch.setattr(orchestrator.AiMatchingClient, "call_json", fake_call_json)
    client = TestClient(app)
    response = client.post(
        "/api/process",
        files={
            "file": (
                workbook_path.name,
                workbook_path.read_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={
            "column_mapping": json.dumps({"输出-价格列": "价格"}, ensure_ascii=False),
            "header_row": "1",
            "only_match_rows_with_value": "false",
            "defer_matching": "true",
            "matching_mode": "ai_enhanced",
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["summary"]["matching_status"] == "pending"
    assert payload["summary"]["total_data_rows"] == 1
    assert payload["ai_matching"]["ai_stage"] == "awaiting_analysis"

    analyzed = client.post(f"/api/process/{payload['job_id']}/ai-matching/analyze", json={})
    assert analyzed.status_code == 200
    analyzed_state = analyzed.json()["ai_matching"]
    assert analyzed_state["ai_stage"] == "awaiting_confirmation"
    assert analyzed_state["metrics"]["model_calls"] == 2
    assert analyzed_state["metrics"]["model_rounds"] == 2

    confirmed = client.post(
        f"/api/process/{payload['job_id']}/ai-matching/confirm",
        json={
            "proposal_ids": [],
            "mapping": analyzed_state["model_result"]["mapping"],
            "expected_revision": analyzed_state["operation_revision"],
        },
    )
    assert confirmed.status_code == 200
    assert confirmed.json()["ai_matching"]["ai_stage"] == "confirmed"

    blocked = client.post("/api/process/batch-match", json={"job_id": payload["job_id"]})
    assert blocked.status_code == 200, blocked.text


def test_ai_candidate_confirmation_reads_trusted_price_and_keeps_evidence(tmp_path, monkeypatch):
    workbook_path = tmp_path / "candidate-input.xlsx"
    _write_candidate_workbook(workbook_path)

    async def fake_call_json(self, messages, *, started_at):
        self.calls += 1
        prompt = json.loads(messages[-1]["content"])
        mapping = {
            "要素1": "项目内容（要素1）",
            "要素2": "专业类别（要素2）",
            "要素4": "比例尺（要素4）",
            "要素5": "复杂程度（要素5）",
            "单位": "计量单位（单位）",
        }
        if prompt.get("phase") == "structure":
            return {"mapping": mapping, "proposals": []}
        record_id = prompt["structure"]["records"][0]["record_id"]
        return {
            "mapping": mapping,
            "proposals": [
                {
                    "record_id": record_id,
                    "field": "price",
                    "candidate_id": "kb-2",
                    "evidence_cells": ["表2!A2", "表2!B2", "表2!C2", "表2!D2", "表2!E2"],
                    "missing_conditions": [],
                    "decision": "accepted",
                }
            ],
        }

    monkeypatch.setattr(orchestrator.AiMatchingClient, "call_json", fake_call_json)
    client = TestClient(app)
    response = client.post(
        "/api/process",
        files={
            "file": (
                workbook_path.name,
                workbook_path.read_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={
            "column_mapping": json.dumps({"输出-价格列": "价格"}, ensure_ascii=False),
            "header_row": "1",
            "only_match_rows_with_value": "false",
            "defer_matching": "true",
            "matching_mode": "ai_enhanced",
        },
    )

    assert response.status_code == 200, response.text
    payload = response.json()
    analyzed = client.post(f"/api/process/{payload['job_id']}/ai-matching/analyze", json={})
    assert analyzed.status_code == 200, analyzed.text
    analyzed_state = analyzed.json()["ai_matching"]
    price_proposals = [item for item in analyzed_state["proposals"] if item.get("field") == "price"]
    assert price_proposals
    assert price_proposals[0]["candidate_id"] == "kb-2"
    assert price_proposals[0]["evidence_cells"] == ["表2!A2", "表2!B2", "表2!C2", "表2!D2", "表2!E2"]

    confirmed = client.post(
        f"/api/process/{payload['job_id']}/ai-matching/confirm",
        json={
            "proposal_ids": [price_proposals[0]["proposal_id"]],
            "mapping": analyzed_state["model_result"]["mapping"],
            "expected_revision": analyzed_state["operation_revision"],
        },
    )
    assert confirmed.status_code == 200, confirmed.text
    batch = client.post("/api/process/batch-match", json={"job_id": payload["job_id"]})
    assert batch.status_code == 200, batch.text
    assert batch.json()["summary"]["matching_status"] == "completed"

    output = client.get(batch.json()["downloads"]["excel"])
    assert output.status_code == 200
    from openpyxl import load_workbook

    workbook = load_workbook(__import__("io").BytesIO(output.content), data_only=False)
    try:
        sheet = workbook["表2"]
        headers = [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]
        price_column = headers.index("价格") + 1
        assert sheet.cell(2, price_column).value == 8033
    finally:
        workbook.close()
