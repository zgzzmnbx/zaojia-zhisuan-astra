from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from app.adjustment_rules import AdjustmentEngine
from app.ai_matching.contracts import empty_ai_state
from app.ai_matching.orchestrator import _validate_model_output, confirm_proposals, prepare_run
from app.ai_matching.store import save_ai_state
from app.ai_matching.workbook_reader import read_workbook_structure
from app.fill_engine import FillEngine
from app.knowledge_base import KnowledgeBase
from app.professional_skills import SkillRuntimeContext


def _write_standard_unknown_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "表2"
    sheet.append(["要素1", "要素2", "要素3", "要素4", "要素5", "单位", "价格", "数量"])
    sheet.append(["知识库不存在的业务描述", "地质测绘", "", "", "", "km2", None, 1])
    workbook.save(path)
    workbook.close()


def _write_semantic_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "表2"
    sheet.append(["要素1", "要素2", "要素3", "要素4", "要素5", "单位", "价格", "实物工作费调整系数", "技术工作费调整系数"])
    sheet.append(["现场原始描述", "现场专业描述", "", "", "", "点", None, None, None])
    workbook.save(path)
    workbook.close()


def _minimal_context(tmp_path: Path) -> SkillRuntimeContext:
    return SkillRuntimeContext(
        skill_id="repair-test-skill",
        skill_version="1.0.0",
        manifest_hash="repair-test-manifest",
        input_profile={"extensions": [".xlsx"]},
        processor_id="survey-measurement-v1",
        knowledge_base_path=tmp_path / "missing-kb.xlsx",
        rule_assets={
            "technicalRules": (tmp_path / "technical.xlsx",),
            "physicalRules": (tmp_path / "physical.xlsx",),
        },
        risk_profile={"experiencePool": tmp_path / "missing-pool.xlsx"},
        knowledge_sources=(),
        report_template_path=tmp_path / "missing-template.docx",
        validation_profile={},
        capabilities={"pricing": True},
    )


def _knowledge_base() -> KnowledgeBase:
    return KnowledgeBase(
        [
            {
                "要素1": "标准描述",
                "要素2": "标准专业",
                "要素3": "",
                "要素4": "",
                "要素5": "",
                "单位": "点",
                "基价": 1234,
                "【经验数】实物工作费调整系数": 0.9,
                "【经验数解释】-实物工作费调整系数": "历史经验",
                "【经验数】技术工作费调整系数": 1.1,
                "【经验数解释】-技术工作费调整系数": "历史经验",
                "_excel_row": 2,
            }
        ]
    )


def test_repair_rejects_rejected_missing_condition_and_unverifiable_evidence(tmp_path):
    job_dir = tmp_path / "job"
    state = empty_ai_state(requested_mode="ai_enhanced", effective_mode="ai_enhanced")
    state.update(
        {
            "ai_stage": "awaiting_confirmation",
            "operation_revision": 2,
            "proposals": [
                {
                    "proposal_id": "aip_rejected",
                    "record_id": "表2!2",
                    "field": "price",
                    "candidate_id": "kb-2",
                    "decision": "rejected",
                    "evidence_cells": [],
                    "missing_conditions": ["专业条件不充分"],
                    "status": "needs_review",
                }
            ],
        }
    )
    save_ai_state(job_dir, state)

    with pytest.raises(ValueError, match="不可确认|拒绝|证据|条件"):
        confirm_proposals(job_dir, proposal_ids=["aip_rejected"], expected_revision=2)

    after = state = __import__("app.ai_matching.store", fromlist=["load_ai_state"]).load_ai_state(job_dir)
    assert after["ai_stage"] == "awaiting_confirmation"
    assert after.get("confirmed_candidates", []) == []
    assert after.get("confirmed_proposal_ids", []) == []


def test_repair_standard_headers_still_route_unknown_business_row_to_ai(tmp_path):
    workbook_path = tmp_path / "standard-unknown.xlsx"
    _write_standard_unknown_workbook(workbook_path)

    state = prepare_run(
        tmp_path / "job",
        workbook_path,
        requested_mode="ai_enhanced",
        options={"ai_matching": {"maxInputRows": 25}},
        context=_minimal_context(tmp_path),
    )

    assert state["structure_kind"] == "standard"
    assert state["ai_stage"] == "awaiting_analysis"
    assert state["coverage"]["business_rows"] == 1
    assert state["pending_ai_review_count"] >= 1


def test_repair_unmapped_rows_keep_real_source_coordinates(tmp_path):
    workbook_path = tmp_path / "unmapped.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "费用明细"
    sheet.append(["业务描述", "专业类别", "计量单位", "数量"])
    sheet.append(["现场原始描述", "地质测绘", "点", 1])
    workbook.save(workbook_path)
    workbook.close()

    structure = read_workbook_structure(workbook_path, header_row=1, only_match_rows_with_value=True, match_value_filter_field="数量")

    assert structure["records"][0]["raw_source_cells"]
    assert structure["records"][0]["raw_source_cells"]["A2"]["cell"] == "A2"


def test_repair_cross_record_candidate_is_not_accepted_as_current_record_candidate():
    structure = {
        "records": [
            {
                "record_id": "表2!2",
                "values": {"要素1": "甲", "单位": "点"},
                "source_cells": {"要素1": {"sheet": "表2", "cell": "A2", "value": "甲"}},
                "candidate_fields": {"price": []},
            },
            {
                "record_id": "表2!3",
                "values": {"要素1": "乙", "单位": "点"},
                "source_cells": {"要素1": {"sheet": "表2", "cell": "A3", "value": "乙"}},
                "candidate_fields": {
                    "price": [
                        {
                            "id": "kb-2",
                            "field": "price",
                            "value": 200,
                            "source": "knowledge_base",
                            "candidate_key": {"要素1": "乙", "单位": "点"},
                        }
                    ]
                },
            },
        ],
        "sheets": [{"headers": ["要素1", "单位"]}],
    }

    _, proposals = _validate_model_output(
        {
            "mapping": {},
            "proposals": [
                {
                    "record_id": "表2!2",
                    "field": "price",
                    "candidate_id": "kb-2",
                    "evidence_cells": ["表2!A2"],
                    "missing_conditions": [],
                    "decision": "accepted",
                }
            ],
        },
        structure,
        [],
    )

    assert len(proposals) == 1
    assert proposals[0]["can_confirm"] is False
    assert "当前记录" in proposals[0]["validation_error"]


def test_repair_confirmed_semantic_price_candidate_does_not_require_raw_key_match(tmp_path):
    workbook_path = tmp_path / "semantic.xlsx"
    _write_semantic_workbook(workbook_path)
    output_path = tmp_path / "output.xlsx"
    engine = FillEngine(_knowledge_base(), AdjustmentEngine([], []))

    engine.fill_workbook(
        workbook_path,
        output_path,
        confirmed_candidates=[
            {
                "proposal_id": "aip_price",
                "record_id": "表2!2",
                "sheet_name": "表2",
                "excel_row": 2,
                "field": "price",
                "candidate_id": "kb-2",
                "semantic_values": {
                    "要素1": "标准描述",
                    "要素2": "标准专业",
                    "要素3": "",
                    "要素4": "",
                    "要素5": "",
                    "单位": "点",
                },
            }
        ],
    )

    from openpyxl import load_workbook

    workbook = load_workbook(output_path, data_only=True)
    try:
        assert workbook["表2"]["G2"].value == 1234
    finally:
        workbook.close()


def test_repair_confirmed_physical_and_technical_candidates_are_consumed(tmp_path):
    workbook_path = tmp_path / "coefficients.xlsx"
    _write_semantic_workbook(workbook_path)
    output_path = tmp_path / "output-coefficients.xlsx"
    engine = FillEngine(_knowledge_base(), AdjustmentEngine([], []))

    engine.fill_workbook(
        workbook_path,
        output_path,
        column_mapping={
            "输出-价格列": "价格",
            "输出-实物工作费调整系数": "实物工作费调整系数",
            "输出-技术工作费调整系数": "技术工作费调整系数",
        },
        confirmed_candidates=[
            {
                "proposal_id": "aip_physical",
                "record_id": "表2!2",
                "sheet_name": "表2",
                "excel_row": 2,
                "field": "physical_adjustment",
                "candidate_id": "kb-2-physical",
                "semantic_values": {"要素1": "标准描述", "要素2": "标准专业", "单位": "点"},
            },
            {
                "proposal_id": "aip_technical",
                "record_id": "表2!2",
                "sheet_name": "表2",
                "excel_row": 2,
                "field": "technical_adjustment",
                "candidate_id": "kb-2-technical",
                "semantic_values": {"要素1": "标准描述", "要素2": "标准专业", "单位": "点"},
            },
        ],
    )

    from openpyxl import load_workbook

    workbook = load_workbook(output_path, data_only=True)
    try:
        headers = [workbook["表2"].cell(1, column).value for column in range(1, workbook["表2"].max_column + 1)]
        physical_column = headers.index("实物工作费调整系数") + 1
        technical_column = headers.index("技术工作费调整系数") + 1
        assert workbook["表2"].cell(2, physical_column).value == 0.9
        assert workbook["表2"].cell(2, technical_column).value == 1.1
    finally:
        workbook.close()
