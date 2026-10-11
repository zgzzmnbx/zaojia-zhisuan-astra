"""Exercise the frozen report pipeline through Astra's real provider adapter."""
import asyncio
import hashlib
import io
import json
import sys
import threading
from pathlib import Path

import pytest
from docx import Document
from fastapi import HTTPException
from openpyxl import Workbook

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'03-Src'))
import server
import model_adapter
import report_markdown
from app.llm import LlmConfig

ANSWER='\n'.join(['一、本次处理结论','本次已完成匹配，结果待人工核验。','二、主要风险概览','建议复核未匹配事项。','三、重点复核建议','请核查输入资料和适用规则。','四、规则依据提示','仅依据所提供的资料，不裁决最终单价。','五、后续处理建议','完成复核后形成正式成果。'])

@pytest.fixture
def report(tmp_path,monkeypatch):
    job='f'*32;folder=tmp_path/job;folder.mkdir()
    book=Workbook();book.active.title='费用汇总';book.active.append(['项目','费用（万元）']);book.active.append(['合计（含税）','=940+4.72']);book.save(folder/'output.xlsx');book.close()
    doc=Document();doc.add_paragraph('一、项目概况');doc.add_paragraph('必须保留的项目事实')
    table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='费用';table.cell(0,1).text='万元';table.cell(1,0).text='含税合计';table.cell(1,1).text='944.72'
    doc.add_paragraph('五、其他需要注意的事项');doc.add_paragraph('造价智算匹配报告');doc.add_paragraph('原机械匹配说明')
    doc.save(folder/'report.docx')
    (folder/'report.md').write_text('原机械匹配摘要\n\n## 价格识别日志\n原日志',encoding='utf-8')
    (folder/'大模型风险报告-【codex】.md').write_text('原风险正文',encoding='utf-8')
    state={'input_filename':'真实报告专项.xlsx','output_report':'report.docx','output_excel':'output.xlsx'}
    monkeypatch.setattr(server.engine,'RUNTIME_DIR',tmp_path)
    monkeypatch.setattr(server.paths,'RUNTIME_DIR',tmp_path)
    monkeypatch.setattr(server.engine,'_load_process_state',lambda _:state)
    monkeypatch.setattr(server.engine,'_job_skill_runtime_context',lambda *a:object())
    monkeypatch.setattr(server.engine,'_require_skill_capability',lambda *a:None)
    monkeypatch.setattr(server.engine,'_find_output_excel_files',lambda _: [folder/'output.xlsx'])
    monkeypatch.setattr(server.engine,'_find_report_files',lambda _,suffix:[folder/('report'+suffix)])
    monkeypatch.setattr(server.engine,'_build_risk_report_knowledge_evidence',lambda *a:('核实的测试规则依据',[{'title':'测试规则'}]))
    events=[]
    monkeypatch.setattr(server.engine,'_record_llm_request',lambda config,**kw:events.append(kw['status']))
    monkeypatch.setattr(server.engine,'_sync_business_task_from_job',lambda *a,**kw:{})
    monkeypatch.setattr(server.engine,'publish_sentinel_event',lambda *a:events.append('published'))
    monkeypatch.setattr(server.engine,'_attach_job_skill',lambda data,*a:data)
    monkeypatch.setattr(LlmConfig,'resolved_api_key',lambda self:'test-key')
    return job,folder,state,events

def transport(monkeypatch,content,finish='stop'):
    calls=[]
    def request(req,timeout):
        calls.append(json.loads(req.data))
        return io.BytesIO(json.dumps({'choices':[{'message':{'content':content},'finish_reason':finish}]}).encode())
    monkeypatch.setattr(model_adapter,'urlopen',request)
    return calls

@pytest.mark.parametrize('content,finish',[('','stop'),(None,'stop'),('部分正文','length'),('受限正文','content_filter'),([], 'stop'),('未确认结束的正文',None)])
def test_failed_model_cannot_write_facts_only_report_or_mark_success(report,monkeypatch,content,finish):
    job,folder,_,events=report;transport(monkeypatch,content,finish)
    files=[folder/'report.docx',folder/'output.xlsx',folder/'report.md',folder/'大模型风险报告-【codex】.md']
    before={p:hashlib.sha256(p.read_bytes()).hexdigest() for p in files}
    with pytest.raises(HTTPException) as error:server._risk_report_sync(job,'deepseek','deepseek-flash','https://api.deepseek.com')
    assert error.value.status_code==400
    assert events==['failed']
    assert not (folder/'risk-report.md').exists()
    assert before=={p:hashlib.sha256(p.read_bytes()).hexdigest() for p in files}
    assert model_adapter._report_request.get() is None

@pytest.mark.parametrize('provider,model',[('deepseek','deepseek-flash'),('siliconflow','selected-model')])
def test_model_text_goes_only_to_original_section_and_excel_is_unchanged(report,monkeypatch,provider,model):
    job,folder,_,events=report;calls=transport(monkeypatch,ANSWER)
    before=hashlib.sha256((folder/'output.xlsx').read_bytes()).hexdigest()
    original=[p.text for p in Document(folder/'report.docx').paragraphs]
    result=server._risk_report_sync(job,provider,model,'https://test.invalid')
    assert result['generated_by_model'] is True and result['risk_report_mode']=='model_generated'
    assert result['model_provider']==provider and result['model']==model
    assert result['risk_report'].startswith(ANSWER)
    paragraphs=[p.text for p in Document(folder/'report.docx').paragraphs]
    start=paragraphs.index('五、其他需要注意的事项')
    assert paragraphs[start+1]=='大模型风险提示'
    assert paragraphs[start+2:start+2+len(ANSWER.splitlines())]==ANSWER.splitlines()
    assert [p for p in paragraphs if p in original]==original
    assert hashlib.sha256((folder/'output.xlsx').read_bytes()).hexdigest()==before
    assert (folder/'report.md').read_text(encoding='utf-8').startswith('原机械匹配摘要')
    assert events==['success','published']
    markdown=(folder/'risk-report.md').read_text(encoding='utf-8')
    assert result['risk_report_markdown']==markdown
    assert result['markdown_sha256']==hashlib.sha256(markdown.encode()).hexdigest()
    for text in ['必须保留的项目事实','| 含税合计 | 944.72 |','原机械匹配说明',*ANSWER.splitlines()]:
        assert text in markdown
    receipt=json.loads((folder/'risk-report-source.json').read_text(encoding='utf-8'))
    assert receipt['source_word_sha256']==hashlib.sha256((folder/'report.docx').read_bytes()).hexdigest()
    assert (folder/'astra-risk-reports'/(result['markdown_sha256']+'.md')).read_text(encoding='utf-8')==markdown
    prompt=calls[0]['messages'][1]['content']
    for text in ['五、其他需要注意的事项','五、后续处理建议','原机械匹配摘要','核实的测试规则依据']:
        assert text in prompt
    assert '944.72' in prompt and '不代表 Excel 已完成重算验收' in prompt
    assert calls[0]['max_tokens']==4096
    assert result['debug']['max_tokens']==4096
    assert ('thinking' in calls[0])==(provider=='deepseek')

def test_explicit_demo_report_stays_preset_and_never_calls_model(report,monkeypatch):
    job,folder,state,events=report;state['input_filename']='示例【演示】.xlsx'
    def unexpected(*a,**kw):raise AssertionError('demo called model')
    monkeypatch.setattr(model_adapter,'urlopen',unexpected)
    result=server._risk_report_sync(job,'deepseek','deepseek-flash','https://test.invalid')
    assert result['generated_by_model'] is False and result['risk_report_mode']=='demo_preset'
    assert events==['published']

def test_report_wait_does_not_block_health_or_leak_report_context(monkeypatch):
    release=threading.Event()
    monkeypatch.setattr(server.professional_host,'model_preference',lambda:{'provider':'deepseek','model':'deepseek-flash','baseUrl':'https://test.invalid'})
    async def blocked(**kw):
        assert model_adapter._report_request.get() is not None
        assert release.wait(1), 'model blocked the ASGI loop'
        return {'risk_report':'正文'}
    monkeypatch.setattr(server.engine,'generate_risk_report',blocked)
    async def run():
        asyncio.get_running_loop().call_later(.02,release.set)
        result=await server.risk_report('f'*32,provider=None,model=None,base_url=None)
        assert result['generated_by_model'] is True
        assert model_adapter._report_request.get() is None
    asyncio.run(run())

def test_complete_markdown_preserves_real_automatic_numbers_without_changing_word(tmp_path):
    doc=Document();doc.add_paragraph('第一项依据',style='List Number');doc.add_paragraph('第二项依据',style='List Number')
    path=tmp_path/'report.docx';doc.save(path);before=path.read_bytes()
    result=report_markdown.persist(tmp_path,path,{'generated_by_model':True,'model':'test-model'})
    assert '1. 第一项依据' in result['risk_report_markdown']
    assert '2. 第二项依据' in result['risk_report_markdown']
    assert path.read_bytes()==before
