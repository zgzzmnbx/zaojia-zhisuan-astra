import hashlib
import io
import json
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / '03-Src'))
import server
import documents

BASE = '/api/astra/documents'
WORKSPACE = '/api/astra/workspace'

@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(server.paths, 'RUNTIME_DIR', tmp_path)
    with TestClient(server.app) as client:
        yield client

def upload(client, name, content):
    conversation = client.post(WORKSPACE+'/conversations', json={}).json()['id']
    response = client.post(WORKSPACE+'/assets', data={'conversation_id': conversation}, files={'file': (name, content)})
    assert response.status_code == 200, response.text
    asset = response.json()
    return {'kind': 'asset', 'source_id': asset['id'], 'conversation_id': conversation, 'sha256': asset['sha256']}

def test_markdown_conversion_uses_real_core_and_immutable_artifact(client):
    content = '# 费用核对说明\n\n## 一、费用构成\n\n核对金额128.60万元，单位m³，偏差-3.50%。\n\n| 专业 | 金额 |\n|---|---|\n| 勘察 | 86.40 |\n'
    request = upload(client, '费用.md', content.encode())
    reply = client.post(BASE+'/format', json=request)
    assert reply.status_code == 200, reply.text
    output = reply.json()
    params = {'conversation_id': request['conversation_id']}
    raw = client.get(BASE+'/'+output['id']+'/download', params=params).content
    assert hashlib.sha256(raw).hexdigest() == output['sha256']
    doc = Document(io.BytesIO(raw))
    assert '核对金额128.60万元，单位m³，偏差-3.50%。' in [p.text for p in doc.paragraphs]
    assert doc.tables[0].cell(1, 1).text == '86.40'
    checks = client.get(BASE+'/'+output['id']+'/download', params={**params, 'kind': 'checks'}).json()
    assert checks['manifest']['profile_id'] == 'official_body'
    assert checks['manifest']['profile_version'] == '1.4.0'
    assert not any(c['status'] == 'fail' for c in checks['checks'])
    original = client.get(WORKSPACE+'/assets/'+request['source_id'], params=params).content
    assert original == content.encode()
    other = client.post(WORKSPACE+'/conversations', json={}).json()['id']
    assert client.get(BASE+'/'+output['id']+'/download', params={'conversation_id': other}).status_code == 403

def test_docx_preserves_tables_and_source_package(client):
    doc = Document(); doc.add_heading('核对报告', 0); doc.add_paragraph('金额 1521.56 万元。')
    doc.add_table(rows=2, cols=2).cell(1, 1).text = '保留123.45'
    buffer = io.BytesIO(); doc.save(buffer); original = buffer.getvalue()
    request = upload(client, '原文.docx', original)
    reply = client.post(BASE+'/format', json=request)
    assert reply.status_code == 200, reply.text
    raw = client.get(BASE+'/'+reply.json()['id']+'/download', params={'conversation_id': request['conversation_id']}).content
    output = Document(io.BytesIO(raw))
    assert [p.text for p in output.paragraphs] == [p.text for p in doc.paragraphs]
    assert output.tables[0].cell(1, 1).text == '保留123.45'
    assert hashlib.sha256(original).hexdigest() == request['sha256']
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        assert 'word/document.xml' in archive.namelist()

def test_scope_stale_source_and_invalid_document_fail_closed(client):
    request = upload(client, '无效.docx', b'not a docx')
    assert client.post(BASE+'/format', json={**request, 'sha256': 'old'}).status_code == 409
    other = client.post(WORKSPACE+'/conversations', json={}).json()['id']
    assert client.post(BASE+'/format', json={**request, 'conversation_id': other}).status_code == 403
    assert client.post(BASE+'/format', json=request).status_code == 422
    assert client.get(BASE+'/source', params={**request, 'source_id': '../source'}).status_code == 400

def test_missing_engine_does_not_publish(client, monkeypatch):
    request = upload(client, '正文.md', '# 测试'.encode())
    monkeypatch.setattr(documents, 'ROOT', Path('missing-source').resolve())
    assert client.post(BASE+'/format', json=request).status_code == 503


@pytest.mark.parametrize('kind', ['md', 'docx'])
@pytest.mark.parametrize('mode', ['suggested', 'other', 'none', 'manual'])
def test_four_title_choices_share_original_core_and_keep_financial_text(client, kind, mode):
    if kind == 'md':
        content = '# 风险复核报告\n\n备选标题\n\n## 一、费用核验\n\n金额128.60万元，偏离-3.50%。\n'
        data = content.encode()
    else:
        doc = Document(); doc.add_heading('风险复核报告', 0); doc.add_paragraph('备选标题')
        doc.add_heading('一、费用核验', 1); doc.add_paragraph('金额128.60万元，偏离-3.50%。')
        buffer = io.BytesIO(); doc.save(buffer); data = buffer.getvalue()
    request = upload(client, '风险.'+kind, data)
    prepared_reply = client.post(BASE+'/prepare', json=request)
    assert prepared_reply.status_code == 200, prepared_reply.text
    prepared = prepared_reply.json()
    selection = prepared['suggested'] if mode == 'suggested' else next(row['id'] for row in prepared['candidates'] if row['text'] == '备选标题') if mode == 'other' else mode
    options = {'confirmed_title': selection}
    if mode == 'manual':
        options['title'] = '人工确认的风险报告'
    result = client.post(BASE+'/format', json={**request, 'prepared_id': prepared['id'], 'options': options})
    assert result.status_code == 200, result.text
    output = result.json()
    assert '【调整格式】' in output['name']
    query = {'conversation_id': request['conversation_id']}
    checks = client.get(BASE+'/'+output['id']+'/checks', params=query).json()
    assert checks['manifest']['engine_version'] == '2.5.4'
    assert checks['manifest']['profile_version'] == '1.4.0'
    assert not any(row['status'] == 'fail' for row in checks['checks'])
    title = checks['manifest']['document_title']
    assert title == {'suggested': '风险复核报告', 'other': '备选标题', 'none': None, 'manual': '人工确认的风险报告'}[mode]
    raw = client.get(BASE+'/'+output['id']+'/download', params=query).content
    doc = Document(io.BytesIO(raw))
    assert '金额128.60万元，偏离-3.50%。' in [p.text for p in doc.paragraphs]
    assert client.get(WORKSPACE+'/assets/'+request['source_id'], params=query).content == data


def test_prepared_snapshot_cannot_be_reused_for_another_asset_and_strict_options(client):
    first = upload(client, '一.md', '# 第一份'.encode())
    prepared = client.post(BASE+'/prepare', json=first).json()
    second = upload(client, '二.md', '# 第二份'.encode())
    assert client.post(BASE+'/format', json={**second, 'prepared_id': prepared['id']}).status_code == 409
    assert client.post(BASE+'/format', json={**first, 'options': {'add_page_numbers': 'false'}}).status_code == 400
    assert client.post(BASE+'/format', json={**first, 'options': {'roles': {'b1': 'title'}}}).status_code == 400
    assert client.post(BASE+'/format', json={**first, 'options': {'confirmed_title': 'manual', 'title': ''}}).status_code == 422
    assert client.post(BASE+'/format', json={**first, 'options': {'confirmed_title': 'not-a-candidate'}}).status_code == 422


def test_actual_report_uses_owned_job_and_preserves_original_package(client, tmp_path):
    from fastapi import FastAPI
    from workspace import Workspace
    request = upload(client, '辅助.md', b'# input')
    source_file = tmp_path/'risk.docx'
    doc = Document(); doc.add_heading('风险报告', 0); doc.add_paragraph('偏离率15.26%，金额905.36万元。')
    doc.add_table(rows=2, cols=2).cell(1, 1).text = '128.60'
    doc.save(source_file)
    original = source_file.read_bytes()
    job = 'a'*32
    with Workspace(tmp_path/'astra-workspace').connect() as db:
        db.execute('INSERT INTO conversation_jobs VALUES (?,?)', (job, request['conversation_id']))
    app = FastAPI(); app.include_router(documents.documents_router(lambda: tmp_path, lambda _: source_file))
    with TestClient(app) as report_client:
        payload = {'kind':'report', 'source_id':job, 'conversation_id':request['conversation_id']}
        prepared = report_client.post(BASE+'/prepare', json=payload).json()
        result = report_client.post(BASE+'/format', json={**payload, 'sha256':prepared['sha256'], 'prepared_id':prepared['id'], 'options':{'confirmed_title':prepared['suggested']}})
        assert result.status_code == 200, result.text
        assert source_file.read_bytes() == original
        formatted = report_client.get(BASE+'/'+result.json()['id']+'/download', params={'conversation_id':request['conversation_id']}).content
        reread = Document(io.BytesIO(formatted))
        assert [p.text for p in reread.paragraphs] == [p.text for p in doc.paragraphs]
        assert reread.tables[0].cell(1, 1).text == '128.60'
        source_file.write_bytes(original+b'updated')
        assert report_client.post(BASE+'/format', json={**payload,'sha256':prepared['sha256'],'prepared_id':prepared['id']}).status_code == 409


def test_risk_markdown_is_versioned_and_scoped_to_the_same_report_job(client, tmp_path):
    from fastapi import FastAPI
    from workspace import Workspace
    request = upload(client, '辅助.md', b'# input')
    report_file = tmp_path/'original-report.docx'
    Document().save(report_file)
    original = report_file.read_bytes()
    md = tmp_path/'risk-report.md'
    md.write_text('# 工程风险报告\n\n## 五、其他需要注意的事项\n\n金额905.36万元，偏离15.26%。\n\n|专业|金额|\n|---|---|\n|测量|128.60|',encoding='utf-8')
    receipt={'source_word_sha256':hashlib.sha256(original).hexdigest(),'markdown_sha256':hashlib.sha256(md.read_bytes()).hexdigest()}
    (tmp_path/'risk-report-source.json').write_text(json.dumps(receipt),encoding='utf-8')
    job = 'b'*32
    with Workspace(tmp_path/'astra-workspace').connect() as db:
        db.execute('INSERT INTO conversation_jobs VALUES (?,?)',(job,request['conversation_id']))
    app = FastAPI();app.include_router(documents.documents_router(lambda:tmp_path,lambda _:report_file))
    with TestClient(app) as pipeline:
        payload = dict(kind='risk-markdown',source_id=job,conversation_id=request['conversation_id'])
        prepared = pipeline.post(BASE+'/prepare',json=payload).json()
        output = pipeline.post(BASE+'/format',json={**payload,'sha256':prepared['sha256'],'prepared_id':prepared['id'],'options':{'confirmed_title':prepared['suggested']}})
        assert output.status_code == 200,output.text
        raw = pipeline.get(BASE+'/'+output.json()['id']+'/download',params={'conversation_id':request['conversation_id']}).content
        reread = Document(io.BytesIO(raw))
        assert '金额905.36万元，偏离15.26%。' in [p.text for p in reread.paragraphs]
        assert reread.tables[0].cell(1,1).text=='128.60'
        assert report_file.read_bytes()==original
        download=pipeline.get(BASE+'/source/download',params=payload)
        assert download.status_code==200 and download.content==md.read_bytes()
        report_file.write_bytes(original+b'changed')
        assert pipeline.post(BASE+'/prepare',json=payload).status_code==409
        report_file.write_bytes(original)
        md.write_text('# changed',encoding='utf-8')
        assert pipeline.post(BASE+'/prepare',json=payload).status_code==409
        other = client.post(WORKSPACE+'/conversations',json={}).json()['id']
        assert pipeline.post(BASE+'/prepare',json={**payload,'conversation_id':other}).status_code==403


def test_v31_standard_answer_formats_with_original_business_text_and_tables(client):
    fixture = next((documents.ROOT/'03-Src/engine/08-测试与演示/测试-V3.1-260802').glob('*.docx'))
    original = fixture.read_bytes()
    request = upload(client,'V3.1标准答案.docx',original)
    prepared = client.post(BASE+'/prepare',json=request).json()
    reply = client.post(BASE+'/format',json={**request,'prepared_id':prepared['id'],'options':{'confirmed_title':prepared['suggested'] or 'none'}})
    assert reply.status_code==200,reply.text
    result = Document(io.BytesIO(client.get(BASE+'/'+reply.json()['id']+'/download',params={'conversation_id':request['conversation_id']}).content))
    source = Document(io.BytesIO(original))
    assert ''.join(p.text for p in source.paragraphs).replace('\n','')==''.join(p.text for p in result.paragraphs).replace('\n','')
    assert [[[c.text for c in row.cells] for row in t.rows] for t in source.tables]==[[[c.text for c in row.cells] for row in t.rows] for t in result.tables]
    assert fixture.read_bytes()==original


def test_control_price_cover_uses_approved_template_layout_and_separate_body(client):
    from report_cover import package,xml,NS,W,boundary,text
    from copy import deepcopy
    from lxml import etree
    fixture=documents.ROOT/'03-Src/reference-assets/cost-report-cover-v31.docx'
    original=fixture.read_bytes()
    request=upload(client,'招标控制价报告.docx',original)
    prepared=client.post(BASE+'/prepare',json=request).json()
    reply=client.post(BASE+'/format',json={**request,'prepared_id':prepared['id'],'options':{'confirmed_title':'manual','title':'正文复核结论'}})
    assert reply.status_code==200,reply.text
    raw=client.get(BASE+'/'+reply.json()['id']+'/download',params={'conversation_id':request['conversation_id']}).content
    output=Document(io.BytesIO(raw));assert len(output.sections)==2
    assert output.sections[0].top_margin!=output.sections[1].top_margin
    assert '正文复核结论' in [p.text for p in output.paragraphs]
    source_parts=package(original);parts=package(raw)
    before=xml(source_parts['word/document.xml']).find('w:body',NS)
    after=xml(parts['word/document.xml']).find('w:body',NS)
    count=boundary(before);actual=[deepcopy(n) for n in list(after)[:count]]
    for node in actual:
        for section in node.findall('.//w:sectPr',NS):section.getparent().remove(section)
        prop=node.find('w:pPr',NS)
        if prop is not None and len(prop)==0:node.remove(prop)
    assert [etree.tostring(n,method='c14n',exclusive=True) for n in list(before)[:count]]==[etree.tostring(n,method='c14n',exclusive=True) for n in actual]
    # Original cover styles keep their complete XML definitions.
    before_styles=xml(source_parts['word/styles.xml']);after_styles=xml(parts['word/styles.xml'])
    for style in before_styles.findall('w:style',NS):
        result=next(n for n in after_styles.findall('w:style',NS) if n.get(W+'styleId')==style.get(W+'styleId'))
        assert etree.tostring(style,method='c14n',exclusive=True)==etree.tostring(result,method='c14n',exclusive=True)
    checks=client.get(BASE+'/'+reply.json()['id']+'/checks',params={'conversation_id':request['conversation_id']}).json()
    assert any(c['code']=='COST_REPORT_COVER' and c['status']=='pass' for c in checks['checks'])
    assert hashlib.sha256(raw).hexdigest()==checks['manifest']['output_sha256']
    assert fixture.read_bytes()==original


def test_cover_keeps_sdk_implicit_body_font_instead_of_cover_default():
    import sys
    sys.path.insert(0,str(documents.ROOT/'03-Src/vendor/gongwen-core'))
    from gongwen_core import format_document,FormatRequest,FormatOptions,MarkdownSource
    from report_cover import preserve_cover,package,xml,NS,W
    original=(documents.ROOT/'03-Src/reference-assets/cost-report-cover-v31.docx').read_bytes()
    markdown='## 一、项目基本情况\n\n正文金额128.60万元。\n\n1. 依据905.36万元。\n\n|金额|\n|---|\n|128.60|'
    generated=format_document(FormatRequest(source=MarkdownSource(markdown),profile_id='official_body',profile_version='1.4.0',options=FormatOptions(confirmed_title='none')))
    raw,checks=preserve_cover(generated.docx_bytes,original)
    result=Document(io.BytesIO(raw));paragraph=next(p for p in result.paragraphs if p.text.startswith('正文金额'))
    assert paragraph.style.style_id=='AstraBody_Normal'
    assert paragraph.style.font.size.pt==16
    assert paragraph.style.paragraph_format.line_spacing.pt==28
    styles=xml(package(raw)['word/styles.xml'])
    assert styles.xpath('./w:style[@w:default="1" and @w:type="paragraph"]/@w:styleId',namespaces=NS)==['a']
    assert checks['cover_unchanged'] and checks['body_content_preserved']
    assert result.tables[-1].cell(1,0).text=='128.60'


def test_cover_only_removes_the_duplicate_model_chapter_wrapper():
    import sys
    sys.path.insert(0,str(documents.ROOT/'03-Src/vendor/gongwen-core'))
    from gongwen_core import format_document,FormatRequest,FormatOptions,MarkdownSource
    from report_cover import preserve_cover
    original=(documents.ROOT/'03-Src/reference-assets/cost-report-cover-v31.docx').read_bytes()
    markdown='# 风险复核\n\n## 一、项目基本情况\n\n金额128.60万元。\n\n## 五、其他需要注意的事项\n\n### 大模型风险提示\n\n## 五、其他需要注意的事项\n\n#### 一、本次处理结论\n\n金额905.36万元，偏离15.26%。'
    generated=format_document(FormatRequest(source=MarkdownSource(markdown),profile_id='official_body',profile_version='1.4.0',options=FormatOptions(clean_extra_spaces=False,convert_chinese_punctuation=False)))
    raw,checks=preserve_cover(generated.docx_bytes,original)
    result=Document(io.BytesIO(raw));lines=[p.text for p in result.paragraphs]
    assert lines.count('五、其他需要注意的事项')==1
    assert '金额128.60万元。' in lines and '金额905.36万元，偏离15.26%。' in lines
    assert '一、本次处理结论' in lines and '大模型风险提示' in lines
    assert checks['duplicate_chapter_heading_removed']==1
    assert checks['body_content_preserved'] and checks['body_text_preserved'] is False


def test_model_subheadings_only_change_inside_the_fifth_chapter():
    from report_cover import normalize_risk_subheadings
    original='''#### 一、本次处理结论

## 五、其他需要注意的事项

### 大模型风险提示

```
#### 一、本次处理结论
```

    一、本次处理结论

> 二、主要风险概览

#### 一、本次处理结论

金额905.36万元，偏离15.26%。

## 二、主要风险概览

三、重点复核建议

| 标题 | 金额 |
| --- | --- |
| 三、重点复核建议 | 128.60 |

### 造价智算匹配报告

#### 四、规则依据提示
'''
    normalized,changes=normalize_risk_subheadings(original)
    assert len(changes)==3
    assert '### （一）本次处理结论' in normalized and '### （二）主要风险概览' in normalized
    assert '### （三）重点复核建议' in normalized
    assert normalized.count('#### 一、本次处理结论')==2
    assert '    一、本次处理结论' in normalized and '> 二、主要风险概览' in normalized
    assert '#### 四、规则依据提示' in normalized
    assert '金额905.36万元，偏离15.26%。' in normalized
    assert '| 三、重点复核建议 | 128.60 |' in normalized
    assert normalize_risk_subheadings(normalized)==(normalized,[])


def test_cost_risk_markdown_uses_sdk_second_level_headings_without_rewriting_sources(client,tmp_path):
    from fastapi import FastAPI
    from workspace import Workspace
    from report_cover import RISK_SUBHEADINGS
    request=upload(client,'辅助.md',b'# input')
    original=(documents.ROOT/'03-Src/reference-assets/cost-report-cover-v31.docx').read_bytes()
    report_file=tmp_path/'report.docx';report_file.write_bytes(original)
    markdown='# 风险复核\n\n## 一、项目基本情况\n\n金额128.60万元。\n\n## 五、其他需要注意的事项\n\n### 大模型风险提示\n\n'
    markdown+='\n\n'.join('#### '+number+'、'+title+'\n\n保留事实905.36万元，偏离15.26%。' for number,title in zip('一二三四五',RISK_SUBHEADINGS))
    markdown+='\n\n### 造价智算匹配报告\n\n|金额|\n|---|\n|128.60|\n'
    md=tmp_path/'risk-report.md';md.write_text(markdown,encoding='utf-8')
    receipt={'source_word_sha256':hashlib.sha256(original).hexdigest(),'markdown_sha256':hashlib.sha256(md.read_bytes()).hexdigest()}
    (tmp_path/'risk-report-source.json').write_text(json.dumps(receipt),encoding='utf-8')
    job='c'*32
    with Workspace(tmp_path/'astra-workspace').connect() as db:
        db.execute('INSERT INTO conversation_jobs VALUES (?,?)',(job,request['conversation_id']))
    app=FastAPI();app.include_router(documents.documents_router(lambda:tmp_path,lambda _:report_file))
    with TestClient(app) as pipeline:
        payload=dict(kind='risk-markdown',source_id=job,conversation_id=request['conversation_id'])
        prepared=pipeline.post(BASE+'/prepare',json=payload).json()
        reply=pipeline.post(BASE+'/format',json={**payload,'sha256':prepared['sha256'],'prepared_id':prepared['id'],'options':{'confirmed_title':'none'}})
        assert reply.status_code==200,reply.text
        raw=pipeline.get(BASE+'/'+reply.json()['id']+'/download',params={'conversation_id':request['conversation_id']}).content
        output=Document(io.BytesIO(raw))
        for number,title in zip('一二三四五',RISK_SUBHEADINGS):
            paragraph=next(p for p in output.paragraphs if p.text=='（'+number+'）'+title)
            assert paragraph.style.style_id=='AstraBody_Heading2'
            assert paragraph.style.font.size.pt==16
            assert paragraph.style.paragraph_format.line_spacing.pt==28
            assert paragraph.style.font.name=='方正楷体简体'
        assert '一、项目基本情况' in [p.text for p in output.paragraphs]
        assert sum(p.text=='保留事实905.36万元，偏离15.26%。' for p in output.paragraphs)==5
        assert output.tables[-1].cell(1,0).text=='128.60'
        checks=pipeline.get(BASE+'/'+reply.json()['id']+'/checks',params={'conversation_id':request['conversation_id']}).json()
        assert len(checks['manifest']['cost_report_subheadings'])==5
        assert checks['manifest']['source_sha256']==receipt['markdown_sha256']
        assert checks['manifest']['cost_report_cover']['cover_unchanged']
        assert any(c['code']=='MODEL_SUBHEADING_LEVEL' and c['status']=='pass' for c in checks['checks'])
        assert report_file.read_bytes()==original and md.read_text(encoding='utf-8')==markdown
