import hashlib
import importlib.util
import io
import json
import shutil
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook,load_workbook

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'03-Src'))
import server
import professional_host
from app import settlement_audit_api

def test_markdown_model_key_reads_only_explicit_section(tmp_path,monkeypatch):
    note=tmp_path/'credentials.md'
    selected='sk-selected000000000000'
    other='sk-other000000000000000'
    note.write_text('### Another\n'+other+'\n### DeepSeek官方\n'+selected+'\n### Other\n'+other,encoding='utf-8')
    settings={'models':{'deepseek':{'enabled':True,'key_file':str(note)+'#DeepSeek官方'}}}
    monkeypatch.setattr(professional_host,'config',lambda:settings)
    assert professional_host.model_key()==selected
    settings['models']['deepseek']['key_file']=str(note)+'#Missing'
    assert professional_host.model_key()==''
    settings['models']['deepseek']['key_file']=str(note)+'#DeepSeek官方'
    note.write_text('### DeepSeek官方\n'+selected+'\n'+other,encoding='utf-8')
    assert professional_host.model_key()==''

def test_missing_model_is_explicit_and_configured_question_uses_model(client,monkeypatch):
    async def search(payload):return {'results':[{'snippet':'source'}],'evidence_found':True}
    monkeypatch.setattr(server.engine,'knowledge_search',search)
    response=client.post('/api/astra/knowledge/ask',json={'question':'非预设问题'})
    assert response.status_code==200
    data=response.json();assert data['model_configured'] is False
    assert '未配置' in data['answer'] and '尚未生成答案' in data['answer']
    calls=[]
    async def answer(payload):
        calls.append(payload);return {'answer':'模型答案','sources':[],'answer_mode':'model_answer'}
    monkeypatch.setattr(professional_host,'model_key',lambda provider:'test-key')
    monkeypatch.setattr(server.engine,'knowledge_ask',answer)
    response=client.post('/api/astra/knowledge/ask',json={'question':'另一个非预设问题'})
    assert response.json()['answer']=='模型答案' and len(calls)==1

def legacy_helpers(name):
    spec=importlib.util.spec_from_file_location(name,server.ENGINE/'backend/tests'/f'{name}.py')
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module);return module

@pytest.fixture
def client(tmp_path,monkeypatch):
    for module in (server.paths,server.engine):monkeypatch.setattr(module,'RUNTIME_DIR',tmp_path)
    monkeypatch.setattr(server.engine,'DEFAULT_KNOWLEDGE_MEMORY_DB_PATH',tmp_path/'memory.sqlite3')
    monkeypatch.setattr(settlement_audit_api,'SETTLEMENT_AUDIT_RUNTIME_DIR',tmp_path/'settlement')
    monkeypatch.setattr(professional_host,'CONFIG',tmp_path/'connections.json')
    with TestClient(server.app) as c:yield c

def test_all_professional_read_contracts(client):
    for endpoint in ['professional-skills/management','professional-skills/open-format','experience-pool/field-preferences',
        'workload-capture/field-preferences','workload-capture/target-field-preferences','experience-warnings/settings',
        'settlement-audit/profile','projects/map','projects/unlocated','projects/spatial-candidates','ui-preferences','preview-column-preferences']:
        r=client.get('/api/'+endpoint);assert r.status_code==200,(endpoint,r.text)

def test_settlement_sample_produces_three_readable_artifacts(client):
    sample=client.get('/api/settlement-audit/sample');assert sample.status_code==200
    r=client.post('/api/settlement-audit/review',files={'file':('sample.xlsx',sample.content)})
    assert r.status_code==200,r.text
    result=r.json();assert result['risks']
    for kind,url in result['downloads'].items():
        artifact=client.get(url);assert artifact.status_code==200
        assert len(artifact.content)>100
        if kind=='excel':
            book=load_workbook(io.BytesIO(artifact.content));assert book.sheetnames;book.close()

def test_workload_same_names_preserve_two_inputs_and_downloads(client,tmp_path):
    helper=legacy_helpers('test_workload_capture')
    source=tmp_path/'source.xlsx';target=tmp_path/'target.xlsx'
    helper._write_source(source);helper._write_target(target)
    r=client.post('/api/workload-capture/run',files={'workload_file':('same.xlsx',source.read_bytes()),'target_file':('same.xlsx',target.read_bytes())},data={
        'source_sheet_configs':json.dumps(helper._source_configs()),'target_sheet_configs':json.dumps(helper._target_configs()),'selected_fields':json.dumps(['数量(信息抓取)'])})
    assert r.status_code==200,r.text
    data=r.json();assert data['summary']['filled_rows']==1
    for url in data['downloads'].values():assert client.get(url).status_code==200
    book=load_workbook(io.BytesIO(client.get(data['downloads']['target']).content));assert book.active['F2'].value==26;book.close()

def test_memory_manual_review_audit_and_project_scope(client):
    helper=legacy_helpers('test_knowledge_memory')
    r=client.post('/api/knowledge-memory/candidates',json=helper.candidate_payload())
    assert r.status_code==200,r.text
    item=r.json()['item'];identifier=item['id'];key=item['project_key']
    assert item['status']!='confirmed'
    base=f'/api/knowledge-memory/items/{identifier}'
    assert client.get(base,params={'project_key':'different-project'}).status_code==404
    for action in ['submit','confirm','revoke']:
        response=client.post(base+'/'+action,json={'project_key':key,'actor':'迁移验收','actor_role':'reviewer','reason':'独立测试资料核验'})
        assert response.status_code==200,response.text
    assert response.json()['item']['status']=='revoked'
    audit=client.get(base+'/audit',params={'project_key':key});assert audit.status_code==200
    assert 'revoke' in audit.text

def test_experience_import_writes_only_working_pool(client,tmp_path,monkeypatch):
    frozen=server.ENGINE/'05-经验池-预警数据/【经验池】-管勘智算-【codex】.xlsx'
    before=hashlib.sha256(frozen.read_bytes()).hexdigest()
    pool=tmp_path/'pool.xlsx';shutil.copy2(frozen,pool);monkeypatch.setattr(server.engine,'DEFAULT_EXPERIENCE_POOL_PATH',pool)
    book=Workbook();ws=book.active;ws.title='历史数据'
    ws.append(['要素1','要素2','要素3','要素4','单位','基价','工程量','实物工作费调整系数','技术工作费调整系数'])
    ws.append(['控制测量','首级控制测量','GPS测量C级','中等','个',123,26,0.6,0.22])
    out=io.BytesIO();book.save(out);book.close()
    inspect=client.post('/api/experience-pool/inspect',files={'file':('history.xlsx',out.getvalue())});assert inspect.status_code==200
    sheets=inspect.json()['sheets'];configs=[{**s,'enabled':True,'column_mapping':s['suggested_mapping']} for s in sheets]
    r=client.post('/api/experience-pool/import',files={'file':('history.xlsx',out.getvalue())},data={'sheet_configs':json.dumps(configs),'selected_fields':json.dumps(['要素1','要素2','要素3','要素4','单位','基价','工程量','实物工作费调整系数','技术工作费调整系数'])})
    assert r.status_code==200,r.text
    assert hashlib.sha256(frozen.read_bytes()).hexdigest()==before
    assert hashlib.sha256(pool.read_bytes()).hexdigest()!=before

def test_model_credentials_do_not_inherit_original_env_or_leak(client,tmp_path,monkeypatch):
    monkeypatch.setenv('DEEPSEEK_API_KEY','original-project-secret')
    from app.llm import LlmConfig
    assert LlmConfig().resolved_api_key()==''
    key=tmp_path/'model.key';key.write_text('astra-test-secret')
    reply=client.put('/api/astra/connections',json={'models':{'deepseek':{'enabled':True,'key_file':str(key)}}})
    assert reply.status_code==200 and reply.json()['models']['deepseek']['configured']
    assert 'astra-test-secret' not in reply.text
    assert LlmConfig().resolved_api_key()=='astra-test-secret'
    assert client.post('/api/collaboration/feishu-webhook/test',json={}).status_code==409

def test_candidate_adoption_persists_pending_state(client,tmp_path,monkeypatch):
    job='d'*32;(tmp_path/job).mkdir();(tmp_path/job/'process-state.json').write_text('{}')
    async def adopt(_):return {'job_id':job,'formula_recalculated':False}
    monkeypatch.setattr(server.engine,'confirm_fill_assist',adopt)
    r=client.post('/api/fill-assist/confirm',json={'job_id':job});assert r.status_code==200
    assert server.read_edit_state(tmp_path/job)


def test_rejected_connection_form_is_not_partially_saved(client,tmp_path):
    previous={'models':{'deepseek':{'enabled':False,'key_file':''}}}
    professional_host.CONFIG.write_text(json.dumps(previous),encoding='utf-8')
    response=client.put('/api/astra/connections',json={'external_enabled':True,'weknora':{'base_url':'https://external.example'}})
    assert response.status_code==400
    assert professional_host.config()==previous
    (tmp_path/('d'*32)).mkdir()
    (tmp_path/('d'*32)/'process-state.json').write_text('{}')
    assert client.post('/api/collaboration/external-dispatch/web-review',json={'job_id':'d'*32}).status_code==409
    assert client.get('/api/settlement-audit/profile').json()['max_file_mb']==30

def test_explicit_platform_connection_preserves_settings_and_uses_selected_platform(client,monkeypatch):
    from app import feishu_app_bot
    previous={'models':{'deepseek':{'enabled':False,'key_file':'keep.key'}},'collaboration_file':'keep.json','external_enabled':False}
    professional_host.CONFIG.write_text(json.dumps(previous),encoding='utf-8')
    monkeypatch.setattr(feishu_app_bot,'load_credentials',lambda _: {'app_id':'test-id','app_secret':'test-secret'})
    monkeypatch.setattr(feishu_app_bot,'credential_configuration_issue',lambda *_: '')
    monkeypatch.setattr(feishu_app_bot,'credential_profiles',lambda:[{'profile_id':'default'},{'profile_id':'weact_cost'}])
    monkeypatch.setattr(feishu_app_bot,'_read_enabled_control',lambda _:True)
    controls=[]
    monkeypatch.setattr(feishu_app_bot,'save_bot_enabled',lambda enabled,profile:controls.append((enabled,profile)))
    calls=[]
    def update(payload):
        calls.append(payload)
        assert professional_host.config()['external_enabled']
        return {'profiles':[]}
    monkeypatch.setattr(server.engine,'update_feishu_app_bot_settings',update)
    assert client.post('/api/astra/connections/platforms/weact_cost/connection',json={'enabled':True}).status_code==200
    assert calls==[{'profile_id':'weact_cost','enabled':True}]
    assert controls==[(False,'default'),(False,'weact_cost')]
    assert professional_host.config()=={**previous,'external_enabled':True}
    assert client.post('/api/astra/connections/platforms/default/connection',json={'enabled':False}).status_code==200
    assert calls[-1]=={'profile_id':'default','enabled':False}

def test_connection_rejects_missing_credentials_and_rolls_back_start_failure(client,monkeypatch):
    from app import feishu_app_bot
    previous={'external_enabled':False,'models':{'deepseek':{'key_file':'keep.key'}}}
    professional_host.CONFIG.write_text(json.dumps(previous),encoding='utf-8')
    endpoint='/api/astra/connections/platforms/default/connection'
    monkeypatch.setattr(feishu_app_bot,'load_credentials',lambda _: {})
    assert client.post(endpoint,json={'enabled':True}).status_code==409
    assert professional_host.config()==previous
    assert client.post(endpoint,json={'enabled':'yes'}).status_code==400
    assert client.post('/api/astra/connections/platforms/other/connection',json={'enabled':True}).status_code==400
    monkeypatch.setattr(feishu_app_bot,'load_credentials',lambda _: {'app_id':'test-id','app_secret':'test-secret'})
    monkeypatch.setattr(feishu_app_bot,'credential_configuration_issue',lambda *_: '')
    def fail(_):raise server.HTTPException(409,'start failed')
    monkeypatch.setattr(server.engine,'update_feishu_app_bot_settings',fail)
    controls=[]
    monkeypatch.setattr(feishu_app_bot,'credential_profiles',lambda:[{'profile_id':'default'}])
    monkeypatch.setattr(feishu_app_bot,'_read_enabled_control',lambda _:False)
    monkeypatch.setattr(feishu_app_bot,'save_bot_enabled',lambda enabled,profile:controls.append((enabled,profile)))
    assert client.post(endpoint,json={'enabled':True}).status_code==409
    assert professional_host.config()==previous
    assert controls==[(False,'default'),(False,'default')]

def test_supervisor_uses_astra_launcher_and_deduplicates_process(monkeypatch,tmp_path):
    import subprocess
    from contextlib import contextmanager
    from app import feishu_app_bot
    monkeypatch.setattr(feishu_app_bot,'RUNTIME_ROOT',tmp_path)
    @contextmanager
    def lock(_):yield True
    monkeypatch.setattr(feishu_app_bot,'_profile_start_lock',lock)
    running=False
    monkeypatch.setattr(feishu_app_bot,'_pid_is_running',lambda _:running)
    calls=[]
    class Process:pid=123
    def launch(args,**kwargs):calls.append((args,kwargs));return Process()
    monkeypatch.setattr(subprocess,'Popen',launch)
    assert feishu_app_bot.start_local_bot_supervisor()
    assert Path(calls[0][0][1])==professional_host.ROOT/'tools/run_collaboration_supervisor.py'
    assert calls[0][1]['cwd']==professional_host.ROOT
    running=True
    assert feishu_app_bot.start_local_bot_supervisor()
    assert len(calls)==1


def test_stopped_receiver_is_not_reported_as_starting(monkeypatch,tmp_path):
    from app import feishu_app_bot
    monkeypatch.setattr(feishu_app_bot,'DB_PATH',tmp_path/'tasks.sqlite3')
    monkeypatch.setattr(feishu_app_bot,'credential_profiles',lambda:[{'profile_id':'default'}])
    monkeypatch.setattr(feishu_app_bot,'load_credentials',lambda _:{'app_id':'id','app_secret':'test-only'})
    monkeypatch.setattr(feishu_app_bot,'credential_configuration_issue',lambda *_:'')
    monkeypatch.setattr(feishu_app_bot,'is_bot_enabled',lambda _:True)
    monkeypatch.setattr(feishu_app_bot,'bot_process_running',lambda _:False)
    profile=feishu_app_bot.bot_status(tmp_path/'tasks.sqlite3')['profiles'][0]
    assert profile['status']=='stopped' and not profile['running']
