import hashlib
import json
import sys
import time
from dataclasses import replace
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'03-Src'))
import server
import professional_host
import business_settings
from app.schemas import FillSummary

@pytest.fixture
def local(tmp_path,monkeypatch):
    for module in (server.engine,server.paths):
        monkeypatch.setattr(module,'RUNTIME_DIR',tmp_path)
        for name in ['DEFAULT_PROJECT_LEDGER_DB_PATH','DEFAULT_BUSINESS_TASK_DB_PATH','DEFAULT_LLM_USAGE_DB_PATH','DEFAULT_KNOWLEDGE_MEMORY_DB_PATH']:
            monkeypatch.setattr(module,name,tmp_path/(name+'.sqlite3'))
    monkeypatch.setattr(professional_host,'CONFIG',tmp_path/'connections.json')
    with TestClient(server.app) as client:yield client,tmp_path

@pytest.fixture
def row(local,monkeypatch):
    client,root=local
    kb=root/'kb.xlsx';book=Workbook();sheet=book.active
    headers=['要素1','要素2','要素3','要素4','要素5','单位','基价','实物工作费调整系数','技术工作费调整系数']
    sheet.append(headers);sheet.append(['控制测量','GPS测量E级','','简单','','点',3000,1,0.22]);book.save(kb);book.close()
    job='c'*32;folder=root/job;folder.mkdir();output=folder/'output.xlsx'
    book=Workbook();sheet=book.active;sheet.title='表2';sheet.append(headers);sheet.append(['控制测量','GPS测量E级','','中等','','点','待复核',1.5,0.33]);book.save(output);book.close()
    summary=FillSummary(total_data_rows=1,price_column='基价',filled_rows=0,matched_rows=0,unchanged_rows=0,review_rows=1,conflict_rows=0,output_excel='output.xlsx',output_report='',report_text='',table_preview={'sheet_name':'表2','header_row':1,'headers':headers,'rows':[['控制测量','GPS测量E级','','中等','','点','待复核',1.5,0.33]],'row_numbers':[2]})
    state={'input_filename':'input.xlsx','input_excel':'','output_excel':'output.xlsx','output_report':'','summary':summary.to_dict(),'artifact_version':1}
    (folder/'process-state.json').write_text(json.dumps(state,ensure_ascii=False),encoding='utf-8')
    original=server.engine._job_skill_runtime_context
    def runtime(*args,**kwargs):return replace(original(*args,**kwargs),knowledge_base_path=kb,risk_profile={'experiencePool':root/'missing.xlsx'})
    monkeypatch.setattr(server.engine,'_job_skill_runtime_context',runtime)
    conversation=client.post('/api/astra/workspace/conversations',json={'title':'专项验收'}).json()['id']
    bound=client.put(f'/api/astra/workspace/conversations/{conversation}',json={'revision':0,'write_token':'migration-test','title':'专项验收','state':{'job_id':job}})
    assert bound.status_code==200,bound.text
    return client,root,output,{'job_id':job,'conversation_id':conversation,'context':'test-context','sheet_name':'表2','row_number':2,'target_header':'基价'}

def candidate(row):
    client,_,_,payload=row
    response=client.post('/api/astra/row-assist/candidates',json=payload)
    assert response.status_code==200,response.text
    return response.json()

def test_row_candidates_are_readonly_and_confirm_only_one_price_field(row):
    client,root,output,_=row
    before=hashlib.sha256(output.read_bytes()).hexdigest();data=candidate(row)
    assert data['candidates'][0]['value']==3000
    assert hashlib.sha256(output.read_bytes()).hexdigest()==before
    reply=client.post('/api/astra/row-assist/confirm',json={**data,'candidate_id':data['candidates'][0]['id'],'note':'人工核对相似计价记录'})
    assert reply.status_code==200,reply.text
    assert reply.json()['needs_recalculate'] is True
    book=load_workbook(output);assert book['表2']['G2'].value==3000;assert book['表2']['H2'].value==1.5;assert book['表2']['I2'].value==0.33;assert '辅助填价人工确认' in book['表2']['G2'].comment.text;book.close()
    assert client.post('/api/astra/row-assist/confirm',json={**data,'candidate_id':data['candidates'][0]['id']}).status_code==409

def test_manual_change_invalidates_the_old_candidate_and_retains_pending_recalculation(row):
    client,root,output,payload=row;data=candidate(row)
    reply=client.post('/api/preview/cell',json={'job_id':payload['job_id'],'sheet_name':'表2','row_number':2,'column_number':7,'value':999,'edit_note':'人工修改','recalculate':False})
    assert reply.status_code==200,reply.text
    rejected=client.post('/api/astra/row-assist/confirm',json={**data,'candidate_id':data['candidates'][0]['id']})
    assert rejected.status_code==409,rejected.text
    book=load_workbook(output);assert book['表2']['G2'].value==999;book.close()
    assert client.get('/api/astra/runs/'+payload['job_id']).json()['needs_recalculate'] is True

def test_tampered_row_field_and_foreign_conversation_are_rejected(row):
    client,_,output,payload=row;data=candidate(row)
    for key,value in [('row_number',3),('column_number',8),('job_id','d'*32),('artifact_version',2)]:
        assert client.post('/api/astra/row-assist/confirm',json={**data,'row_context':{**data['row_context'],key:value},'candidate_id':data['candidates'][0]['id']}).status_code==409
    other=client.post('/api/astra/workspace/conversations',json={'title':'其他对话'}).json()['id']
    assert client.post('/api/astra/row-assist/candidates',json={**payload,'conversation_id':other}).status_code==409
    coefficient=client.post('/api/astra/row-assist/candidates',json={**payload,'target_header':'实物工作费调整系数'})
    assert coefficient.status_code==200 and coefficient.json()['row_context']['column_number']==8

def test_business_settings_conflict_cannot_overwrite_newer_choices(local):
    client,root=local;current=client.get('/api/astra/business-settings').json()
    changed=client.put('/api/astra/business-settings',json={**current,'matchingMode':'ai_enhanced','onlyMatchRowsWithValue':False})
    assert changed.status_code==200,changed.text
    assert client.put('/api/astra/business-settings',json=current).status_code==409
    assert client.get('/api/astra/business-settings').json()==changed.json()
    assert client.put('/api/astra/business-settings',json={**changed.json(),'matchValueFilterField':'基价'}).status_code==400

def test_selected_provider_risk_report_and_ai_snapshot_do_not_require_deepseek(local,monkeypatch):
    client,root=local;config={'provider':'siliconflow','model':'selected-model','baseUrl':'https://api.siliconflow.cn','temperature':0.7,'maxTokens':2048}
    professional_host.CONFIG.write_text(json.dumps({'model':config}),encoding='utf-8')
    job='e'*32;folder=root/job;folder.mkdir();(folder/'process-state.json').write_text(json.dumps({'input_filename':'真实项目.xlsx'}),encoding='utf-8')
    providers=[]
    monkeypatch.setattr(professional_host,'model_key',lambda provider='deepseek':providers.append(provider) or ('test-only' if provider=='siliconflow' else ''))
    async def report(**kwargs):return {'provider':kwargs['provider'],'model':kwargs['model'],'base_url':kwargs['base_url']}
    monkeypatch.setattr(server.engine,'generate_risk_report',report)
    response=client.post('/api/risk-report',data={'job_id':job})
    assert response.status_code==200,response.text
    assert response.json()['provider']=='siliconflow' and providers==['siliconflow']
    settings=server.engine._load_ai_matching_settings()
    from app.ai_matching.client import AiMatchingClient
    captured=AiMatchingClient.from_options({'ai_matching':settings})
    assert captured.public_config()['provider']=='siliconflow'
    professional_host.CONFIG.write_text(json.dumps({'model':{**config,'provider':'deepseek'}}),encoding='utf-8')
    assert captured.public_config()['provider']=='siliconflow'
    calls=[]
    async def transport(self,selected,messages,*,deadline):
        calls.append((selected.provider,selected.model,selected.base_url,selected.temperature,selected.max_tokens));return '{"verified":true}'
    monkeypatch.setattr(AiMatchingClient,'_request_async',transport)
    import asyncio,time
    assert asyncio.run(captured.call_json([{'role':'user','content':'contract test'}],started_at=time.monotonic()))=={'verified':True}
    assert calls==[('siliconflow','selected-model','https://api.siliconflow.cn',0.7,2048)]


def test_host_identity_matches_loaded_source_and_has_no_claimed_model_connection(local):
    client,_=local;health=client.get('/api/health').json()
    assert health['version']==server.app.version
    assert health['host_sources']['server.py']==server.HOST_SOURCE_SHA256
    assert health['workspace']==str(server.ROOT)
    assert {'row-assist-v1','batch-match-recovery-v1','business-settings-v1'}<=set(health['capabilities'])
    assert health['model_connected'] is False


def test_row_snapshot_uses_business_task_version_and_attempt(row):
    client,root,output,payload=row
    folder=root/payload['job_id'];file=folder/'process-state.json';state=json.loads(file.read_text(encoding='utf-8'));state.pop('artifact_version');state['business_task']={'task_id':'task-a','artifact_version':7};state['project_tracking']={'run_id':'run-a'};file.write_text(json.dumps(state),encoding='utf-8');(folder/'batch-match-status.json').write_text(json.dumps({'status':'completed','attempt':'attempt-a'}),encoding='utf-8')
    data=candidate(row);assert data['row_context']['artifact_version']==7;assert data['row_context']['task_id']=='task-a';assert data['row_context']['run_id']=='run-a';assert data['row_context']['attempt_id']=='attempt-a'
    state['business_task']['artifact_version']=8;file.write_text(json.dumps(state),encoding='utf-8');assert client.post('/api/astra/row-assist/confirm',json={**data,'candidate_id':data['candidates'][0]['id']}).status_code==409


def test_late_row_answer_is_rejected_if_artifact_changes(row,monkeypatch):
    client,root,output,payload=row;data=candidate(row)
    async def late_answer(payload):
        book=load_workbook(output);book['表2']['G2']=999;book.save(output);book.close()
        return {'answer':'old row answer','generated_by_model':False}
    monkeypatch.setattr(server,'knowledge_answer',late_answer)
    reply=client.post('/api/astra/row-assist/ask',json={**data,'question':'核对当前行'})
    assert reply.status_code==409,reply.text


def test_price_details_real_samples_and_rule_ranking_are_readonly(row,monkeypatch):
    client,root,output,payload=row
    pool=root/'pool.xlsx';book=Workbook();sheet=book.active
    headers=['要素1','要素2','要素3','要素4','要素5','单位','基价','来源文件','来源sheet','来源行']
    sheet.append(headers)
    for i,value in enumerate([3100,2900,3000]):sheet.append(['控制测量','GPS测量E级','','中等','','点',value,'原始资料.xlsx','表2',i+2])
    sheet.append(['其他工程','','','','','m',99999,'其他资料','表1',8])
    book.save(pool);book.close()
    original=server.engine._job_skill_runtime_context
    monkeypatch.setattr(server.engine,'_job_skill_runtime_context',lambda *a,**kw:replace(original(*a,**kw),risk_profile={'experiencePool':pool}))
    data=candidate(row);before=hashlib.sha256(output.read_bytes()).hexdigest()
    reply=client.post('/api/astra/row-assist/details',json=data)
    assert reply.status_code==200,reply.text
    body=reply.json();experience=body['experience']
    assert experience['total']==3 and experience['distribution']['mean']==3000
    assert [s['value'] for s in experience['items']]==[2900,3000,3100]
    assert sum(b['count'] for b in experience['distribution']['bins'])==3
    assert body['rules']['items'][0]['value']==3000
    assert hashlib.sha256(output.read_bytes()).hexdigest()==before


def test_rule_value_change_rejects_old_confirmation_and_details(row):
    client,root,_,payload=row;data=candidate(row)
    book=load_workbook(root/'kb.xlsx');book.active['G2']=1234;book.save(root/'kb.xlsx');book.close()
    for endpoint in ['details','confirm','ask']:
        reply=client.post('/api/astra/row-assist/'+endpoint,json={**data,'candidate_id':data['candidates'][0]['id']})
        assert reply.status_code==409,reply.text


def test_price_header_row_is_rejected_by_backend(row):
    client,_,_,payload=row
    reply=client.post('/api/astra/row-assist/candidates',json={**payload,'row_number':1})
    assert reply.status_code==400 and '表头' in reply.json()['detail']


def test_price_distribution_counts_all_samples_and_handles_flat_and_empty():
    import row_assist
    assert row_assist.distribution([])=={'count':0,'bins':[]}
    assert row_assist.distribution([0,0,None])['bins']==[{'min':0,'max':0,'count':2}]
    stats=row_assist.distribution(list(range(1000)))
    assert stats['count']==1000 and stats['median']==499.5 and sum(b['count'] for b in stats['bins'])==1000


@pytest.mark.parametrize('metric,column,other,value',[('实物工作费调整系数','H2','I2',1.25),('技术工作费调整系数','I2','H2',0.35)])
def test_coefficient_uses_its_own_candidates_and_only_writes_selected_field(row,metric,column,other,value):
    client,root,output,payload=row
    book=load_workbook(root/'kb.xlsx');sheet=book.active;sheet['D2']='中等';sheet['H1']='【经验数】实物工作费调整系数';sheet['I1']='【经验数】技术工作费调整系数';sheet['H2']=1.25;sheet['I2']=0.35;book.save(root/'kb.xlsx');book.close()
    book=load_workbook(output);book['表2'][column[0]+'1']=metric+'-测试';book.save(output);book.close()
    data=client.post('/api/astra/row-assist/candidates',json={**payload,'target_header':metric}).json()
    candidate=next(c for c in data['candidates'] if c['value']==value)
    assert candidate['id'].endswith('physical' if column=='H2' else 'technical')
    assert all(c['value']!=3000 for c in data['candidates'])
    details=client.post('/api/astra/row-assist/details',json=data)
    assert details.status_code==200 and details.json()['rules']['distribution']['max']<10
    before=load_workbook(output);other_value=before['表2'][other].value;price=before['表2']['G2'].value;before.close()
    confirmed=client.post('/api/astra/row-assist/confirm',json={**data,'candidate_id':candidate['id'],'note':'逐字段确认系数'})
    assert confirmed.status_code==200,confirmed.text
    book=load_workbook(output);assert book['表2'][column].value==value;assert book['表2'][other].value==other_value;assert book['表2']['G2'].value==price;book.close()


def test_model_analysis_does_not_lock_reading_details(row,monkeypatch):
    import asyncio,threading
    from concurrent.futures import ThreadPoolExecutor
    client,_,_,_=row;data=candidate(row);started=threading.Event();release=threading.Event()
    async def pending_answer(payload):
        started.set();await asyncio.to_thread(release.wait,5)
        return {'answer':'当前候选核对完成','generated_by_model':True}
    monkeypatch.setattr(server,'knowledge_answer',pending_answer)
    with ThreadPoolExecutor(max_workers=2) as executor:
        future=executor.submit(client.post,'/api/astra/row-assist/ask',json=data)
        try:
            assert started.wait(3)
            reply=client.post('/api/astra/row-assist/details',json=data)
            assert reply.status_code==200,reply.text
            assert not future.done()
        finally:release.set()
        assert future.result().status_code==200


def test_row_candidate_normalization_matches_batch_engine(row):
    client,root,output,payload=row
    book=load_workbook(output);book['表2']['C2']='成图比例';book.save(output);book.close()
    data=candidate(row)
    assert data['context']['row']['要素3']==''
    assert data['candidates'][0]['similarity']>60


def test_model_shortlist_matches_visible_distinct_prices_and_retains_support():
    import row_assist
    items=[{'id':str(i),'value':v} for i,v in enumerate([5,5,0,8,9])]
    result=row_assist.recommendation_shortlist(items)
    assert [c['value'] for c in result]==[5,0,8]
    assert len(result[0]['supporting_sources'])==2
